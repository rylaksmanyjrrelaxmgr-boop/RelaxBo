#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
handlers_callback.py - معالج الأزرار (v9.4.19)
=====================================================================
🆕 v9.4.19 — إصلاحات v9.4.18:
  - updates_channel_btn: فحص نهائي لصحة URL قبل تمريره لـ Telegram
    (كان يُرسل روابط غير صالحة → BadRequest صامت → زر معطّل)
  - updates_channel_btn في قائمة `known` لمعالجة parameterized
  - توحيد escape للنصوص

✅ v9.4.18 — إضافة زر قناة التحديثات:
  - _show_updates_channel: يعرض قناة التحديثات برابط قابل للنقر
  - معالج updates_channel_btn في handle()
  - يدعم: @username، معرّف رقمي، رابط

✅ v9.4.17 — ترجمة كاملة
✅ v9.4.16 — إصلاح أزرار التحليلات
=====================================================================
"""

import asyncio
import logging
import json
import time
import shutil
import os
import re
import html as _html
import random
from datetime import timedelta
from pathlib import Path
from typing import Optional, Dict, Tuple, Any, Set, List

from telegram import (
    Update, InlineKeyboardButton, InlineKeyboardMarkup,
    LabeledPrice, ChatPermissions
)
from telegram.ext import ContextTypes
from telegram.error import BadRequest, RetryAfter, Forbidden

from config import CONFIG, PATHS
from database import DB, TimeUtils, internal_cache

try:
    from database_analytics import color_emoji
except ImportError:
    def color_emoji(value, thresholds=(0.3, 0.7), inverse=False):
        try:
            v = float(value)
        except (TypeError, ValueError):
            return "⚪"
        low, high = thresholds
        if inverse:
            return "🟢" if v <= low else ("🟡" if v <= high else "🔴")
        return "🟢" if v >= high else ("🟡" if v >= low else "🔴")

try:
    from utils import (
        safe_send, is_authorized_in_group,
        get_text, StateManager, UserState,
        KeyboardFactory, CB, get_ram_usage,
        SmartCache, TranslationManager,
    )
except ImportError:
    from .utils import (
        safe_send, is_authorized_in_group,
        get_text, StateManager, UserState,
        KeyboardFactory, CB, get_ram_usage,
        SmartCache, TranslationManager,
    )

try:
    from utils import PUBLISH_RATE_LIMITER
except ImportError:
    try:
        from utils import RATE_LIMITER as PUBLISH_RATE_LIMITER
    except ImportError:
        class _NullLimiter:
            async def acquire(self):
                return
            async def __aenter__(self):
                return self
            async def __aexit__(self, *a):
                return False
        PUBLISH_RATE_LIMITER = _NullLimiter()

try:
    from cache import user_cache, invalidate_user_cache, posts_cache
except ImportError:
    class _DummyUserCache:
        async def get(self, user_id):
            return None
        async def get_or_load(self, user_id, db):
            try:
                return await db.get_start_data(user_id) or {}
            except Exception:
                return {}
        async def invalidate(self, user_id=None):
            return

    class _DummyPostsCache:
        async def invalidate(self, *a, **k):
            return

    user_cache = _DummyUserCache()
    posts_cache = _DummyPostsCache()

    async def invalidate_user_cache(user_id):
        return

try:
    from .handlers_command import CommandHandlers, _invalidate_force_sub_cache
except ImportError:
    from handlers_command import CommandHandlers, _invalidate_force_sub_cache

logger = logging.getLogger(__name__)

# =====================================================================
# خريطة أسماء أزرار التحليلات
# =====================================================================

_ANALYTICS_ALIASES: Dict[str, str] = {
    "growth_30d_btn": "user_growth",
    "top_channels_btn": "top_channels",
    "publish_stats_btn": "publish_stats",
    "channels_rate_btn": "channels_rate",
    "subscriptions_btn": "subscriptions",
    "pool_live_btn": "pool",
    "slow_queries_btn": "slow",
    "export_excel_btn": "export",
    "user_growth": "user_growth",
    "top_channels": "top_channels",
    "publish_stats": "publish_stats",
    "channels_rate": "channels_rate",
    "subscriptions": "subscriptions",
    "pool": "pool",
    "slow": "slow",
    "export": "export",
}

# =====================================================================
# ثوابت
# =====================================================================

MAX_CAPTION_LENGTH = 1024
MAX_MESSAGE_LENGTH = 4096
MAX_BACKUPS = getattr(CONFIG, "MAX_BACKUPS", 10)
MAX_CONCURRENT_PUBLISH = 2
MAX_PUBLISH_DELAY_SECONDS = 60
MAX_TG_FILE_SIZE = 49 * 1024 * 1024
CALLBACK_MIN_INTERVAL = 0.5
RATE_LIMIT_WINDOW = 60
RATE_LIMIT_CLEANUP_EVERY = 100
ADMIN_PAGE_SIZE = 10
SEC_AUTH_CACHE_TTL = 300
PUBLISH_ACQUIRE_TIMEOUT = 30
SEC_SETTINGS_CACHE_TTL = 5
SEC_STATS_CACHE_TTL = 30
LOG_CHANNEL_MENU_CACHE_TTL = 30
SUCCESS_RATE_GOOD_THRESHOLD = 70
SUCCESS_RATE_WARN_THRESHOLD = 30
DEFAULT_SUCCESS_RATE = 100.0

_BOLD_MD_PATTERN = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)
_VALID_URL_PATTERN = re.compile(r'^https?://[^\s]+$')

try:
    _PRIMARY_OWNER_ID = int(CONFIG.PRIMARY_OWNER_ID)
except (TypeError, ValueError, AttributeError):
    _PRIMARY_OWNER_ID = None

ACTIVE_TASKS: Set[asyncio.Task] = set()
_publish_semaphore = asyncio.Semaphore(MAX_CONCURRENT_PUBLISH)
_sec_auth_cache: Dict[Tuple[int, int], Tuple[bool, float]] = {}
_security_settings_cache: SmartCache = SmartCache(ttl=SEC_SETTINGS_CACHE_TTL, max_size=500)
_security_stats_cache_local: SmartCache = SmartCache(ttl=SEC_STATS_CACHE_TTL, max_size=500)

_CONTEXT_KEYS_TO_CLEAR = (
    'security_chat_id', 'auto_chat', 'adv_chat', 'schedule_ch',
    'ban_chat', 'contest_join', 'channel_page', 'post_page', 'sec_chat',
)
_CANCEL_EXTRA_KEYS = ('pin_msg_id',)

GROUP_NUMBER_EMOJIS = ["1️⃣", "2️⃣", "3️⃣", "4️⃣", "5️⃣", "6️⃣", "7️⃣", "8️⃣", "9️⃣", "🔟"]


# =====================================================================
# ترجمة موحدة
# =====================================================================

async def _trans(key: str, lang: str, default: str = "") -> str:
    if not key:
        return default or ""
    try:
        if lang and lang != 'off':
            text = TranslationManager.get_text(lang, key)
            if text and text != key:
                return text
    except Exception as e:
        logger.debug(f"_trans({key},{lang}) TM: {e}")
    try:
        if lang and lang != 'off':
            text = await get_text(lang, key)
            if text and text != key:
                return text
    except Exception as e:
        logger.debug(f"_trans({key},{lang}) get_text: {e}")
    return default or key


def _fmt(text: str, **kwargs) -> str:
    try:
        return text.format(**kwargs)
    except (KeyError, IndexError):
        return text


# =====================================================================
# دوال مساعدة
# =====================================================================

def _group_number(index: int) -> str:
    if 1 <= index <= len(GROUP_NUMBER_EMOJIS):
        return GROUP_NUMBER_EMOJIS[index - 1]
    return f"{index}."


def _is_primary_owner(user_id: int) -> bool:
    return _PRIMARY_OWNER_ID is not None and user_id == _PRIMARY_OWNER_ID


def _row_to_dict(row) -> Optional[Dict[str, Any]]:
    if row is None:
        return None
    if isinstance(row, dict):
        return row
    try:
        return dict(row)
    except (TypeError, ValueError):
        return None


def _coerce_int(value, default=0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _coerce_float(value, default=0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _safe_str(value, default='?') -> str:
    if value is None:
        return default
    s = str(value)
    return s if s.strip() else default


def _md_to_html(text: str) -> str:
    if not text or '**' not in text:
        return text
    return _BOLD_MD_PATTERN.sub(r"<b>\1</b>", text)


def _log_channel_cache_key(chat_id: int) -> str:
    return f"log_ch_menu_{chat_id}"


async def _invalidate_log_channel_menu_cache(chat_id: int) -> None:
    try:
        await internal_cache.invalidate(_log_channel_cache_key(chat_id))
    except Exception:
        pass
    try:
        await internal_cache.invalidate(f"group_log_{chat_id}")
    except Exception:
        pass


def _set_sec_chat(context, chat_id: int) -> None:
    try:
        context.user_data['sec_chat'] = chat_id
        context.user_data['security_chat_id'] = chat_id
    except Exception:
        pass


def _is_valid_url(url: Optional[str]) -> bool:
    """✅ v9.4.19: فحص نهائي لصلاحية URL قبل تمريره لـ Telegram."""
    if not url:
        return False
    if not isinstance(url, str):
        return False
    url_stripped = url.strip()
    if not url_stripped:
        return False
    if ' ' in url_stripped or '\n' in url_stripped or '\t' in url_stripped:
        return False
    if not _VALID_URL_PATTERN.match(url_stripped):
        return False
    return True


# =====================================================================
# قناة السجل
# =====================================================================

async def _get_log_channel_menu_data(chat_id: int) -> Dict[str, Any]:
    cache_key = _log_channel_cache_key(chat_id)
    cached = await internal_cache.get(cache_key)
    if cached is not None:
        return cached

    current = None
    effective = None

    try:
        current = await DB.get_group_log_channel(chat_id)
    except Exception as e:
        logger.debug(f"get_group_log_channel({chat_id}): {e}")

    if not current:
        try:
            effective = await DB.get_log_channel()
        except Exception as e:
            logger.debug(f"get_log_channel: {e}")

    share_count = 0
    share_names: List[str] = []
    if current:
        try:
            others = await DB.get_groups_sharing_log_channel(
                current, exclude_group_id=chat_id)
            if others:
                share_count = len(others)
                for o in others[:3]:
                    if isinstance(o, dict):
                        name = o.get('chat_name') or f"Group {o.get('chat_id', '?')}"
                    else:
                        name = f"Group {_safe_str(o)}"
                    share_names.append(str(name))
        except Exception as e:
            logger.debug(f"get_groups_sharing({current}): {e}")

    data = {
        'current': current,
        'effective': effective,
        'share_count': share_count,
        'share_names': share_names,
    }
    await internal_cache.set(cache_key, data, ttl=LOG_CHANNEL_MENU_CACHE_TTL)
    return data


async def _safe_answer(query, text=None, show_alert=False) -> bool:
    if not query:
        return False
    try:
        if text:
            await query.answer(text, show_alert=show_alert)
        else:
            await query.answer()
        return True
    except BadRequest:
        return False
    except Exception:
        return False


async def safe_edit(query, text, reply_markup=None, parse_mode=None, bot=None,
                    clear_markup=False) -> bool:
    if not query or not query.message:
        if bot and query and query.from_user:
            try:
                await bot.send_message(
                    chat_id=query.from_user.id,
                    text=text or "...",
                    reply_markup=reply_markup,
                    parse_mode=parse_mode)
                return True
            except Exception:
                pass
        return False

    if not text or not str(text).strip():
        text = "..."
    else:
        text = str(text)

    if clear_markup and reply_markup is None:
        reply_markup = InlineKeyboardMarkup([])

    try:
        await query.edit_message_text(
            text, reply_markup=reply_markup, parse_mode=parse_mode)
        return True
    except BadRequest as e:
        error_msg = str(e).lower()
        if "message is not modified" in error_msg:
            return True
        elif "message is too long" in error_msg:
            chat_id = query.message.chat_id
            try:
                await query.message.delete()
            except Exception:
                pass
            try:
                if bot:
                    await bot.send_message(
                        chat_id=chat_id, text=text,
                        reply_markup=reply_markup, parse_mode=parse_mode)
                    return True
                return False
            except Exception:
                return False
        elif "query is too old" in error_msg or "message to edit not found" in error_msg:
            return False
        elif "can't parse entities" in error_msg or "parse" in error_msg:
            try:
                await query.edit_message_text(
                    text, reply_markup=reply_markup, parse_mode=None)
                return True
            except Exception:
                return False
        elif "message text is empty" in error_msg:
            try:
                await query.edit_message_text("...", reply_markup=reply_markup)
                return True
            except Exception:
                return False
        return False
    except Exception:
        return False


async def safe_delete_message(query_or_message) -> None:
    try:
        if hasattr(query_or_message, 'message') and query_or_message.message:
            await query_or_message.message.delete()
        elif query_or_message:
            await query_or_message.delete()
    except Exception:
        pass


def _mask_id(id_value, prefix=3, suffix=2) -> str:
    if id_value is None:
        return "***"
    s = str(id_value)
    if len(s) <= 5:
        return "***"
    return s[:prefix] + "***" + s[-suffix:]


async def _is_channel_owner(user_id: int, channel_db_id: int) -> bool:
    try:
        return await DB.is_channel_owner(user_id, channel_db_id)
    except Exception:
        return False


async def _is_group_owner(user_id: int, chat_id: int) -> bool:
    try:
        if hasattr(DB, 'is_group_owner'):
            return await DB.is_group_owner(user_id, chat_id)
        if hasattr(DB, 'get_user_groups'):
            groups = await DB.get_user_groups(user_id)
            for g in (groups or []):
                gd = _row_to_dict(g)
                if gd and gd.get('chat_id') == chat_id:
                    return True
        return False
    except Exception as e:
        logger.warning(f"_is_group_owner error: {e}")
        return False


def _clear_context_keys(context, extra_keys=None) -> None:
    for k in _CONTEXT_KEYS_TO_CLEAR:
        context.user_data.pop(k, None)
    if extra_keys:
        for k in extra_keys:
            context.user_data.pop(k, None)


def _ensure_bot_start_time(context) -> None:
    if 'start_time' not in context.bot_data:
        context.bot_data['start_time'] = time.monotonic()


async def _resolve_sec_chat_id(context, data: str) -> Optional[int]:
    parts = data.split(":")
    for part in parts[1:]:
        p = part.strip()
        if p.startswith('-') and p[1:].isdigit():
            return int(p)
    stored = (context.user_data.get('security_chat_id')
              or context.user_data.get('sec_chat'))
    if stored:
        try:
            return int(stored)
        except (TypeError, ValueError):
            return None
    return None


async def _check_sec_auth(context, user_id: int, chat_id: int) -> bool:
    if chat_id is None:
        return False
    key = (user_id, chat_id)
    now = time.monotonic()
    cached = _sec_auth_cache.get(key)
    if cached and now - cached[1] < SEC_AUTH_CACHE_TTL:
        return cached[0]
    try:
        result = await is_authorized_in_group(context.bot, chat_id, user_id)
    except Exception:
        result = False
    _sec_auth_cache[key] = (result, now)
    return result


def _invalidate_sec_auth_cache(chat_id: int = None) -> None:
    if chat_id is None:
        _sec_auth_cache.clear()
    else:
        for k in list(_sec_auth_cache.keys()):
            if k[1] == chat_id:
                del _sec_auth_cache[k]


async def _invalidate_after_channel_change(
    user_id: int,
    channel_db_id: Optional[int] = None,
    invalidate_posts: bool = True,
) -> None:
    keys = [f"start_data_{user_id}", f"user_{user_id}",
            f"user_{user_id}_True", f"user_{user_id}_False",
            f"channels_{user_id}"]
    if channel_db_id is not None:
        keys.append(f"channel_info_{channel_db_id}")
    for k in keys:
        try:
            await internal_cache.invalidate(k)
        except Exception:
            pass
    try:
        await invalidate_user_cache(user_id)
    except Exception as e:
        logger.debug(f"user_cache invalidate: {e}")
    if invalidate_posts and channel_db_id is not None:
        try:
            await posts_cache.invalidate(channel_db_id)
        except Exception as e:
            logger.debug(f"posts_cache invalidate: {e}")


def _format_channel_rate_line(ch: Dict[str, Any]) -> str:
    published = _coerce_int(ch.get('published'), 0)
    failed = _coerce_int(ch.get('failed'), 0)
    total = _coerce_int(ch.get('total'), 0)
    attempted = _coerce_int(ch.get('attempted'), published + failed)
    pending = _coerce_int(ch.get('pending'), max(0, total - attempted))
    success_rate = _coerce_float(ch.get('success_rate'), DEFAULT_SUCCESS_RATE)
    completion_rate = _coerce_float(
        ch.get('completion_rate'),
        (published / total * 100) if total > 0 else 0.0)
    color = color_emoji(success_rate,
        (SUCCESS_RATE_WARN_THRESHOLD, SUCCESS_RATE_GOOD_THRESHOLD))
    name = _html.escape(str(ch.get('name') or '?')[:25])
    line = f"{color} <b>{name}</b>\n"
    line += (f"   🎯 {published}/{attempted}  ❌ {failed}  ({success_rate}%)\n")
    if pending > 0:
        line += f"   📈 {published}/{total}  ⏳ {pending}  ({completion_rate}%)\n"
    line += "\n"
    return line


# =====================================================================
# CallbackHandlers
# =====================================================================

class CallbackHandlers:

    RATE_LIMIT_PER_MINUTE = 30
    PUBLISH_DELAY_SECONDS = 0.2
    PUBLISH_BATCH_SIZE = 10

    @staticmethod
    async def handle(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        query = update.callback_query
        if not query:
            return

        user = query.from_user
        if not user or not user.id:
            await _safe_answer(query, "❌", show_alert=True)
            return

        data = query.data
        if not data:
            await _safe_answer(query)
            return

        user_id = user.id
        now_time = time.monotonic()

        last_cb_key = f"last_cb_{user_id}"
        last_time = context.user_data.get(last_cb_key, 0)
        if now_time - last_time < CALLBACK_MIN_INTERVAL:
            await _safe_answer(query, "⚠️")
            return
        context.user_data[last_cb_key] = now_time

        counter_key = "_cb_counter"
        cb_count = context.bot_data.get(counter_key, 0) + 1
        context.bot_data[counter_key] = cb_count
        if cb_count % RATE_LIMIT_CLEANUP_EVERY == 0:
            CallbackHandlers._cleanup_user_data(context)

        rate_key = f"rate_{user_id}"
        rate_data = context.user_data.get(rate_key)
        if not rate_data or now_time - rate_data.get('reset', 0) > RATE_LIMIT_WINDOW:
            rate_data = {'count': 0, 'reset': now_time}
        rate_data['count'] = rate_data.get('count', 0) + 1
        context.user_data[rate_key] = rate_data
        if rate_data['count'] > CallbackHandlers.RATE_LIMIT_PER_MINUTE:
            await _safe_answer(query, "⚠️", show_alert=True)
            return

        await _safe_answer(query)

        start_time = time.monotonic()
        _ensure_bot_start_time(context)

        try:
            lang = await DB.get_user_language(user_id) or 'ar'
        except Exception:
            lang = 'ar'

        try:
            handled = await CallbackHandlers._handle_parameterized(
                update, context, query, user_id, lang, data)
            if handled:
                return
        except Exception as e:
            logger.error(f"❌ param outer: {e}", exc_info=True)
            try:
                await safe_edit(query, await _trans('error_occurred', lang, "❌"),
                                bot=context.bot)
            except Exception:
                pass
            return

        base_data = data
        if ':' in data:
            parts = data.split(':')
            known = [
                CB.TOGGLE_AUTO, CB.TOGGLE_REC, CB.TRANSLATION, CB.REFERRAL,
                CB.REMINDER, CB.CONTESTS, CB.SUPPORT_TICKET, CB.CH_LIST,
                CB.POST_ADD, CB.POST_PUB, CB.POST_LIST, CB.POST_REC, CB.PUB_ALL,
                CB.GROUPS, CB.ADMIN, CB.SETTINGS, CB.PLANS, CB.INVOICES,
                CB.REF_CLAIM, CB.REF_LIST, CB.CONTEST_WINNERS, CB.DEVELOPER,
                CB.SUBSCRIBE, CB.SUPPORT, CB.LANGUAGE, CB.TRIAL, CB.HELP,
                CB.CANCEL, CB.CHECK_SUB, CB.TRANS_OFF, CB.REM_TOGGLE_SUB,
                CB.REM_TOGGLE_DAILY, CB.REM_TOGGLE_WEEKLY, CB.REM_SET_DAYS,
                CB.ADMIN_LIST_ADMINS, "finish_posts", "gift_plans", "redeem_gift",
                "updates_channel_btn",  # ✅ v9.4.19
            ]
            if parts[0] in known:
                base_data = parts[0]

        try:
            if base_data in (CB.MAIN, CB.BACK):
                StateManager.clear(user_id)
                _clear_context_keys(context)
                ok = await CallbackHandlers._show_main_menu_inline(query, context, user_id)
                if not ok:
                    await CommandHandlers.start(update, context)
                return

            if base_data == CB.CANCEL:
                StateManager.clear(user_id)
                _clear_context_keys(context, extra_keys=list(_CANCEL_EXTRA_KEYS))
                try:
                    await query.edit_message_reply_markup(reply_markup=None)
                except Exception:
                    pass
                return

            if base_data == CB.HELP:
                StateManager.clear(user_id)
                await CommandHandlers.help_command(update, context)
                return

            if base_data == CB.TRIAL:
                if await DB.has_used_trial(user_id):
                    await safe_edit(query,
                        await _trans('trial_used', lang, "❌"),
                        bot=context.bot)
                    return
                days = 0
                try:
                    days = _coerce_int(await DB.activate_trial(user_id), 0)
                finally:
                    try:
                        await DB.invalidate_subscription_cache(user_id)
                    except Exception:
                        pass
                if days == -1:
                    text = await _trans('trial_already_longer', lang, "ℹ️")
                elif days > 0:
                    text = _fmt(await _trans('trial_activated', lang, "✅ {days}"),
                                days=days)
                else:
                    text = await _trans('trial_failed', lang, "❌")
                await safe_edit(query, text, bot=context.bot)
                await invalidate_user_cache(user_id)
                return

            if base_data == CB.DEVELOPER:
                StateManager.clear(user_id)
                await CommandHandlers.developer(update, context)
                return

            if base_data == CB.SUBSCRIBE:
                StateManager.clear(user_id)
                await CommandHandlers.subscribe(update, context)
                return

            if base_data == CB.SUPPORT:
                StateManager.clear(user_id)
                await CommandHandlers.support(update, context)
                return

            if base_data == CB.LANGUAGE:
                StateManager.clear(user_id)
                await CommandHandlers.language(update, context)
                return

            if base_data == CB.CHECK_SUB:
                try:
                    _invalidate_force_sub_cache(user_id)
                except Exception:
                    pass
                StateManager.clear(user_id)
                ok = await CallbackHandlers._show_main_menu_inline(query, context, user_id)
                if not ok:
                    await CommandHandlers.start(update, context)
                return

            if base_data == CB.SETTINGS:
                await CallbackHandlers._render_settings(query, context, user_id, lang)
                return

            if base_data == CB.TOGGLE_AUTO:
                cur = await DB.get_auto_publish_status(user_id)
                await DB.set_auto_publish(user_id, not cur)
                await CallbackHandlers._render_settings(query, context, user_id, lang)
                await invalidate_user_cache(user_id)
                return

            if base_data == CB.TOGGLE_REC:
                cur = await DB.get_auto_recycle_status(user_id)
                await DB.set_auto_recycle(user_id, not cur)
                await CallbackHandlers._render_settings(query, context, user_id, lang)
                await invalidate_user_cache(user_id)
                return

            if base_data == CB.PLANS:
                await safe_edit(query, await _trans('plan_selector', lang, "💎"),
                    reply_markup=KeyboardFactory.build("plans", lang=lang),
                    bot=context.bot)
                return

            if base_data == "gift_plans":
                plans = await DB.get_gift_plans()
                if not plans:
                    await safe_edit(query,
                        await _trans('no_gift_plans', lang, "📭"),
                        bot=context.bot)
                    return
                kb = []
                for p in plans:
                    pd = _row_to_dict(p) or {}
                    days = _safe_str(pd.get('days'), '0')
                    price = _safe_str(pd.get('price'), '0')
                    kb.append([InlineKeyboardButton(
                        f"🎁 {days} - {price} ⭐",
                        callback_data=f"buy_gift:{pd.get('id', 0)}")])
                kb.append([InlineKeyboardButton(
                    KeyboardFactory.get_text("back", lang), callback_data=CB.BACK)])
                await safe_edit(query,
                    await _trans('gift_plans_text', lang, "💎"),
                    reply_markup=InlineKeyboardMarkup(kb), bot=context.bot)
                return

            if base_data == "redeem_gift":
                StateManager.clear(user_id)
                await CommandHandlers.redeem_gift(update, context)
                return

            if base_data == CB.INVOICES:
                invoices = await DB.get_user_invoices(user_id, 10)
                if not invoices:
                    await safe_edit(query,
                        await _trans('no_invoices', lang, "📭"),
                        bot=context.bot)
                    return
                lines = []
                for inv in invoices:
                    inv_d = _row_to_dict(inv) or {}
                    lines.append(f"• #{_safe_str(inv_d.get('number'))} - "
                                 f"{_safe_str(inv_d.get('amount'), '0')} ⭐")
                text = await _trans('invoices_title', lang, "🧾") + "\n\n" + "\n".join(lines)
                await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(KeyboardFactory.get_text("back", lang),
                                           callback_data=CB.BACK)]]),
                    bot=context.bot)
                return

            if base_data == CB.REFERRAL:
                await CallbackHandlers._render_referral(query, context, user_id, lang)
                return

            if base_data == CB.REF_CLAIM:
                days = await DB.claim_referral_reward(user_id)
                if days > 0:
                    text = _fmt(await _trans('referral_claimed', lang,
                                             "✅ {days}!"), days=days)
                else:
                    text = await _trans('no_rewards', lang, "📭")
                await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(KeyboardFactory.get_text("back", lang),
                                           callback_data=CB.REFERRAL)]]),
                    bot=context.bot)
                await invalidate_user_cache(user_id)
                return

            if base_data == CB.REF_LIST:
                refs = await DB.get_referrals_list(user_id)
                if refs:
                    text = await _trans('referrals_list', lang, "📋") + "\n\n" + "\n".join(
                        f"{i}. {_mask_id(r)}" for i, r in enumerate(refs[:20], 1))
                else:
                    text = await _trans('no_data', lang, "📭")
                await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(KeyboardFactory.get_text("back", lang),
                                           callback_data=CB.REFERRAL)]]),
                    bot=context.bot)
                return

            if base_data in (CB.REM_TOGGLE_SUB, CB.REM_TOGGLE_DAILY, CB.REM_TOGGLE_WEEKLY):
                await CallbackHandlers._handle_reminder_toggle(
                    query, context, user_id, lang, base_data)
                return

            if base_data == CB.REMINDER:
                await CallbackHandlers._render_reminder(query, context, user_id, lang)
                return

            if base_data == CB.REM_SET_DAYS:
                StateManager.set(user_id, UserState.WAIT_REM_DAYS)
                await safe_edit(query,
                    await _trans('send_days_btn', lang, "📅"),
                    bot=context.bot)
                return

            if base_data == "rem_lang":
                StateManager.set(user_id, UserState.WAIT_REM_LANG)
                await safe_edit(query,
                    await _trans('choose_language', lang, "🌐"),
                    parse_mode='HTML', bot=context.bot)
                return

            if base_data == CB.TRANSLATION:
                await CallbackHandlers._render_translation_menu(
                    query, context, user_id, lang)
                return

            if base_data == CB.TRANS_OFF:
                await DB.set_user_language(user_id, 'off')
                await safe_edit(query,
                    await _trans('translation_disabled', lang, "✅"),
                    bot=context.bot)
                await invalidate_user_cache(user_id)
                return

            if base_data == CB.CONTESTS:
                StateManager.clear(user_id)
                await CommandHandlers.contests(update, context)
                return

            if base_data == CB.CONTEST_WINNERS:
                winners = await DB.get_contest_winners(10)
                if winners:
                    lines = []
                    for w in winners:
                        wd = _row_to_dict(w) or {}
                        lines.append(f"• {_safe_str(wd.get('title'))} - "
                                     f"{_mask_id(wd.get('winner_id'))}")
                    text = await _trans('winners_title', lang, "🏆") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_winners', lang, "📭")
                await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(KeyboardFactory.get_text("back", lang),
                                           callback_data=CB.BACK)]]),
                    bot=context.bot)
                StateManager.clear(user_id)
                return

            if base_data == CB.SUPPORT_TICKET:
                StateManager.set(user_id, UserState.SUPPORT_MODE)
                await safe_send(context.bot, user_id,
                                await _trans('send_support_message', lang, "📞"))
                return

            if base_data == CB.CH_ADD:
                has_sub = _is_primary_owner(user_id) or await DB.has_active_subscription(user_id)
                if not has_sub:
                    await safe_edit(query,
                        await _trans('subscription_required', lang, "❌"),
                        bot=context.bot)
                    return
                StateManager.set(user_id, UserState.WAIT_CHANNEL)
                await safe_edit(query,
                    await _trans('send_channel_id', lang, "📡"),
                    bot=context.bot)
                return

            if base_data == CB.CH_LIST:
                await CallbackHandlers._show_channel_list(
                    update, context, query, user_id, lang)
                return

            if base_data == CB.POST_ADD:
                await CallbackHandlers._handle_post_add(update, context, query, user_id)
                return

            if base_data == "finish_posts":
                StateManager.clear(user_id)
                await invalidate_user_cache(user_id)
                ok = await CallbackHandlers._show_main_menu_inline(query, context, user_id)
                if not ok:
                    await safe_edit(query,
                        await _trans('post_finish', lang, "✅"),
                        bot=context.bot)
                return

            if base_data == CB.POST_PUB:
                await CallbackHandlers._handle_post_publish(update, context, query, user_id)
                return

            if base_data == CB.POST_LIST:
                await CallbackHandlers._show_post_list(
                    update, context, query, user_id, lang)
                return

            if base_data == CB.POST_REC:
                active = await DB.get_active_channel(user_id)
                if active:
                    count = await DB.reset_posts(user_id, active)
                    await _invalidate_after_channel_change(user_id, active)
                    context.user_data['post_page'] = 0
                    msg = _fmt(await _trans('posts_recycled', lang, "♻️ {count}"),
                               count=count)
                    await safe_send(context.bot, user_id, msg)
                    await CallbackHandlers._show_post_list(
                        update, context, query, user_id, lang)
                else:
                    await safe_edit(query,
                        await _trans('no_active_channel', lang, "❌"),
                        bot=context.bot)
                return

            if base_data == CB.POST_CLEAR:
                active = await DB.get_active_channel(user_id)
                if active:
                    await DB.execute("DELETE FROM posts WHERE channel_db_id=?", (active,))
                    await _invalidate_after_channel_change(user_id, active)
                    context.user_data['post_page'] = 0
                    await safe_send(context.bot, user_id,
                                    await _trans('posts_cleared', lang, "🧹"))
                    await CallbackHandlers._show_post_list(
                        update, context, query, user_id, lang)
                else:
                    await safe_edit(query,
                        await _trans('no_active_channel', lang, "❌"),
                        bot=context.bot)
                return

            if base_data == CB.PUB_ALL:
                await CallbackHandlers._handle_publish_all(update, context, query, user_id)
                return

            if base_data == CB.GROUPS:
                await CallbackHandlers._show_groups_list(
                    update, context, query, user_id, lang)
                return

            # ═════════════════════════════════════════════════════════
            # ✅ v9.4.18: زر قناة التحديثات
            # ═════════════════════════════════════════════════════════
            if base_data == "updates_channel_btn":
                await CallbackHandlers._show_updates_channel(
                    query, context, user_id, lang)
                return

            if base_data == CB.ADMIN:
                if not CONFIG.is_developer(user_id):
                    await safe_edit(query,
                        await _trans('unauthorized', lang, "❌"),
                        bot=context.bot)
                    return
                kb = KeyboardFactory.build("admin_panel", lang=lang)
                await safe_edit(query,
                    await _trans('admin_panel', lang, "👑"),
                    reply_markup=kb, bot=context.bot)
                return

            if data.startswith("log_channel_"):
                await CallbackHandlers._handle_log_channel(
                    update, context, query, user_id, lang)
                return

            if data.startswith("sec_"):
                await CallbackHandlers._handle_security(
                    update, context, query, user_id, lang)
                return

            if data == "admin_analytics":
                if not CONFIG.is_developer(user_id):
                    await safe_edit(query,
                        await _trans('unauthorized', lang, "❌"),
                        bot=context.bot)
                    return
                await CallbackHandlers._show_analytics_menu(
                    query, context, user_id, lang)
                return

            if data == "refresh_btn":
                if not CONFIG.is_developer(user_id):
                    await safe_edit(query,
                        await _trans('unauthorized', lang, "❌"),
                        bot=context.bot)
                    return
                await CallbackHandlers._show_analytics_menu(
                    query, context, user_id, lang)
                return

            if (data.startswith("analytics_") or data in _ANALYTICS_ALIASES):
                if not CONFIG.is_developer(user_id):
                    await safe_edit(query,
                        await _trans('unauthorized', lang, "❌"),
                        bot=context.bot)
                    return
                await CallbackHandlers._handle_analytics(
                    update, context, query, user_id, lang, data)
                return

            if data.startswith("admin_"):
                if CONFIG.is_developer(user_id):
                    await CallbackHandlers._handle_admin(
                        update, context, query, user_id, lang)
                else:
                    await safe_edit(query,
                        await _trans('unauthorized', lang, "❌"),
                        bot=context.bot)
                return

            if data.startswith("auto_reply_") or data.startswith("auto_reply_menu:"):
                await CallbackHandlers._handle_auto_reply(
                    update, context, query, user_id, lang)
                return

            if data.startswith("sched_open:") or data.startswith("sched_"):
                await CallbackHandlers._handle_schedule(
                    update, context, query, user_id)
                return

            if data.startswith("ban_") or data.startswith("act_") or data.startswith("pen_"):
                await CallbackHandlers._handle_advanced_actions(
                    update, context, query, user_id)
                return

            declare_sel = getattr(CB, 'DECLARE_WINNER_SEL', 'declare_winner_sel')
            if data.startswith("contest_") or data.startswith(declare_sel + ":"):
                await CallbackHandlers._handle_contests(
                    update, context, query, user_id)
                return

            if data.startswith("lang_"):
                await CallbackHandlers._handle_language_change(
                    update, context, query, user_id)
                return

            if data == "ch_page_prev":
                context.user_data['channel_page'] = max(
                    0, _coerce_int(context.user_data.get('channel_page'), 0) - 1)
                await CallbackHandlers._show_channel_list(
                    update, context, query, user_id, lang)
                return
            if data == "ch_page_next":
                context.user_data['channel_page'] = _coerce_int(
                    context.user_data.get('channel_page'), 0) + 1
                await CallbackHandlers._show_channel_list(
                    update, context, query, user_id, lang)
                return
            if data == "post_page_prev":
                context.user_data['post_page'] = max(
                    0, _coerce_int(context.user_data.get('post_page'), 0) - 1)
                await CallbackHandlers._show_post_list(
                    update, context, query, user_id, lang)
                return
            if data == "post_page_next":
                context.user_data['post_page'] = _coerce_int(
                    context.user_data.get('post_page'), 0) + 1
                await CallbackHandlers._show_post_list(
                    update, context, query, user_id, lang)
                return

            if data.startswith("adm_ch_page:"):
                page = _coerce_int(data.split(":")[-1], 0)
                context.user_data['adm_ch_page'] = page
                await CallbackHandlers._show_admin_channels(
                    update, context, query, user_id, lang)
                return
            if data.startswith("adm_gr_page:"):
                page = _coerce_int(data.split(":")[-1], 0)
                context.user_data['adm_gr_page'] = page
                await CallbackHandlers._show_admin_groups(
                    update, context, query, user_id, lang)
                return

            if data in ("panel_lock", "panel_unlock", "panel_close"):
                await CallbackHandlers._handle_panel(
                    update, context, query, user_id, data, lang)
                return

            await safe_edit(query,
                await _trans('not_available', lang, "⚠️"),
                bot=context.bot)

        except BadRequest as e:
            if "query is too old" not in str(e).lower():
                logger.error(f"❌ BadRequest: {e}", exc_info=True)
        except Exception as e:
            logger.error(f"❌ Callback error: {e}", exc_info=True)
        finally:
            elapsed = time.monotonic() - start_time
            if elapsed > 1.0:
                logger.warning(f"🐢 Slow button {data[:30]} — {elapsed:.2f}s")

    @staticmethod
    def _cleanup_user_data(context) -> None:
        try:
            now = time.monotonic()
            keys_to_del = []
            for k, v in list(context.user_data.items()):
                if k.startswith('rate_') and isinstance(v, dict):
                    if now - v.get('reset', 0) > RATE_LIMIT_WINDOW * 5:
                        keys_to_del.append(k)
                elif k.startswith('last_cb_'):
                    if isinstance(v, (int, float)) and now - v > 3600:
                        keys_to_del.append(k)
            for k in keys_to_del:
                context.user_data.pop(k, None)
        except Exception:
            pass

    # ═════════════════════════════════════════════════════════════
    # ✅ v9.4.18: قناة التحديثات
    # ✅ v9.4.19: فحص نهائي لصحة URL
    # ═════════════════════════════════════════════════════════════

    @staticmethod
    async def _show_updates_channel(query, context, user_id, lang):
        """
        ✅ v9.4.18: عرض قناة التحديثات للمستخدم برابط قابل للنقر.
        يدعم: @username, رقم صحيح, رابط كامل.

        ✅ v9.4.19: فحص نهائي لصحة URL قبل تمريرها لـ Telegram
        (كان يُرسل روابط غير صالحة → BadRequest صامت → زر معطّل).
        """
        title = await _trans('updates_channel_title', lang, "📢 قناة التحديثات")

        try:
            ch = await DB.get_updates_channel()
        except Exception as e:
            logger.warning(f"get_updates_channel: {e}")
            ch = None

        back_text = KeyboardFactory.get_text("back", lang)

        if not ch:
            text = (f"{title}\n"
                    "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                    + await _trans('no_update_channel', lang,
                                   "📭 لم يتم تعيين قناة تحديثات بعد"))
            kb = InlineKeyboardMarkup([[
                InlineKeyboardButton(back_text, callback_data=CB.BACK)
            ]])
            await safe_edit(query, text, reply_markup=kb,
                            parse_mode='HTML', bot=context.bot)
            return

        ch_str = str(ch).strip()
        url = None
        display = ch_str

        try:
            # 1) @username
            if ch_str.startswith('@'):
                url = f"https://t.me/{ch_str[1:]}"
                display = ch_str

            # 2) رقم سالب → محاولة الحصول على رابط
            elif ch_str.lstrip('-').isdigit():
                cid = int(ch_str)
                try:
                    chat = await context.bot.get_chat(cid)
                    if getattr(chat, 'username', None):
                        url = f"https://t.me/{chat.username}"
                        display = f"@{chat.username}"
                    elif getattr(chat, 'invite_link', None):
                        url = chat.invite_link
                        display = chat.title or ch_str
                    else:
                        try:
                            url = await context.bot.export_chat_invite_link(cid)
                            display = chat.title or ch_str
                        except Exception:
                            url = None
                except Exception as e:
                    logger.debug(f"get_chat {cid}: {e}")
                    url = None

            # 3) رابط كامل
            elif ch_str.startswith(('https://', 'http://')):
                url = ch_str
                display = ch_str
            elif ch_str.startswith(('t.me/', 'telegram.me/')):
                url = f"https://{ch_str}"
                display = url
            else:
                # username بدون @
                url = f"https://t.me/{ch_str}"
                display = f"@{ch_str}"

        except Exception as e:
            logger.debug(f"_show_updates_channel build url: {e}")

        # ✅ v9.4.19: فحص نهائي قبل بناء الزر
        url_is_valid = _is_valid_url(url)
        if url and not url_is_valid:
            logger.warning(
                f"⚠️ v9.4.19: updates_channel URL غير صالح: {url!r} "
                f"(ch_str={ch_str!r}) — سيُعرض بدون زر"
            )

        # ─── بناء النص ───
        text = (
            f"{title}\n"
            "━━━━━━━━━━━━━━━━━━━━━━\n\n"
            f"📌 <b>{_html.escape(display)}</b>"
        )

        # ─── الأزرار ───
        rows = []
        if url_is_valid:
            open_text = await _trans('open_channel_btn', lang, "📢 فتح القناة")
            rows.append([InlineKeyboardButton(open_text, url=url)])
        else:
            text += ("\n\n⚠️ " +
                     await _trans('channel_link_unavailable', lang,
                                  "لا يمكن إنشاء رابط لهذه القناة"))

        rows.append([InlineKeyboardButton(back_text, callback_data=CB.BACK)])

        kb = InlineKeyboardMarkup(rows)
        await safe_edit(query, text, reply_markup=kb,
                        parse_mode='HTML', bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # دوال الأمان
    # =================================================================

    @staticmethod
    async def _get_security_settings_cached(chat_id: int) -> Dict:
        key = f"sec_set_{chat_id}"
        cached = await _security_settings_cache.get(key)
        if cached is not None:
            return cached
        try:
            settings = await DB.get_security_settings(chat_id) or {}
            if not isinstance(settings, dict):
                settings = _row_to_dict(settings) or {}
        except Exception as e:
            logger.error(f"get_security_settings({chat_id}): {e}")
            settings = {}
        await _security_settings_cache.set(key, settings, ttl=SEC_SETTINGS_CACHE_TTL)
        return settings

    @staticmethod
    async def _invalidate_security_settings_cache(chat_id: int) -> None:
        try:
            await _security_settings_cache.delete(f"sec_set_{chat_id}")
        except Exception:
            pass
        try:
            await _security_stats_cache_local.delete(f"sec_stats_{chat_id}")
        except Exception:
            pass

    @staticmethod
    async def _load_stats_and_edit(query, context, chat_id, lang, settings):
        try:
            cache_key = f"sec_stats_{chat_id}"
            cached_stats = await _security_stats_cache_local.get(cache_key)
            if cached_stats is not None:
                stats = cached_stats
            else:
                stats = await KeyboardFactory._get_security_stats(chat_id) or {}
                await _security_stats_cache_local.set(
                    cache_key, stats, ttl=SEC_STATS_CACHE_TTL)
            try:
                text = KeyboardFactory._format_security_text(settings, stats, lang=lang)
            except TypeError:
                text = KeyboardFactory._format_security_text(settings, stats)
            kb = KeyboardFactory.build("security", chat_id=chat_id, lang=lang)
            await safe_edit(query, text, reply_markup=kb,
                            parse_mode='HTML', bot=context.bot)
        except Exception as e:
            logger.debug(f"_load_stats_and_edit({chat_id}): {e}")

    @staticmethod
    async def _render_security_two_phase(query, context, chat_id, lang,
                                          force_refresh_settings=False):
        try:
            if force_refresh_settings:
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
            settings = await CallbackHandlers._get_security_settings_cached(chat_id)
            try:
                text_no_stats = KeyboardFactory._format_security_text(
                    settings, {}, lang=lang)
            except TypeError:
                text_no_stats = KeyboardFactory._format_security_text(settings, {})
            kb = KeyboardFactory.build("security", chat_id=chat_id, lang=lang)
            await safe_edit(query, text_no_stats, reply_markup=kb,
                            parse_mode='HTML', bot=context.bot)
            task = asyncio.create_task(
                CallbackHandlers._load_stats_and_edit(
                    query, context, chat_id, lang, settings))
            ACTIVE_TASKS.add(task)
            task.add_done_callback(ACTIVE_TASKS.discard)
        except Exception as e:
            logger.error(f"_render_security_two_phase: {e}", exc_info=True)

    # ═════════════════════════════════════════════════════════════
    # القائمة الرئيسية
    # =================================================================

    @staticmethod
    async def _show_main_menu_inline(query, context, user_id) -> bool:
        try:
            user_data = None
            try:
                user_data = await user_cache.get(user_id)
            except Exception:
                pass
            if not user_data:
                try:
                    if hasattr(user_cache, 'get_or_load'):
                        user_data = await user_cache.get_or_load(user_id, DB)
                    else:
                        user_data = await DB.get_start_data(user_id) or {}
                except Exception:
                    try:
                        user_data = await DB.get_start_data(user_id) or {}
                    except Exception:
                        user_data = {}
            if not isinstance(user_data, dict):
                user_data = _row_to_dict(user_data) or {}

            lang = user_data.get('language', 'ar') or 'ar'
            channel_info = user_data.get('channel_info')
            unpublished_posts = user_data.get('unpublished_posts', 0)
            groups_count = user_data.get('groups_count', 0)
            has_sub = bool(user_data.get('has_subscription', False))
            auto_publish = bool(user_data.get('auto_publish', True))
            auto_recycle = bool(user_data.get('auto_recycle', True))
            auto_publish_text = "✅" if auto_publish else "❌"
            auto_recycle_text = "✅" if auto_recycle else "❌"

            ch_display = await _trans('no_active_channel', lang, "-")
            if channel_info and isinstance(channel_info, dict):
                raw_name = channel_info.get('channel_name')
                if raw_name:
                    ch_display = _html.escape(str(raw_name))

            if has_sub:
                sub_text = await _trans('subscription_active', lang, "✅")
            else:
                sub_text = await _trans('subscription_inactive', lang, "❌")

            try:
                kb = KeyboardFactory.build("main_menu", lang=lang)
            except Exception:
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.BACK)]])

            if CONFIG.is_developer(user_id):
                admin_text = KeyboardFactory.get_text("admin_panel_btn", lang)
                existing_callbacks = set()
                for row in kb.inline_keyboard:
                    for btn in row:
                        if btn.callback_data:
                            existing_callbacks.add(btn.callback_data)
                if CB.ADMIN not in existing_callbacks:
                    new_rows = list(kb.inline_keyboard)
                    new_rows.append([InlineKeyboardButton(
                        admin_text, callback_data=CB.ADMIN)])
                    kb = InlineKeyboardMarkup(new_rows)

            title = await get_text(
                lang, 'main_menu',
                user_name=f"<code>{user_id}</code>",
                groups_count=groups_count,
                active_channel=ch_display,
                unpublished_posts=unpublished_posts,
                subscription_status=sub_text,
                auto_publish=auto_publish_text,
                auto_recycle=auto_recycle_text,
            )
            title = _md_to_html(title)
            await safe_edit(query, title, reply_markup=kb,
                            parse_mode='HTML', bot=context.bot)
            return True
        except Exception as e:
            logger.error(f"_show_main_menu_inline: {e}", exc_info=True)
            return False

    # ═════════════════════════════════════════════════════════════
    # Parameterized
    # =================================================================

    @staticmethod
    async def _handle_parameterized(update, context, query, user_id, lang, data) -> bool:
        try:
            if data in ("sec_close", "grp_close", "security_close",
                        "back_to_groups", "sec_back"):
                StateManager.clear(user_id)
                _clear_context_keys(context)
                await CallbackHandlers._show_groups_list(
                    update, context, query, user_id, lang)
                return True

            if data.startswith("sec_close:") or data.startswith("grp_close:"):
                StateManager.clear(user_id)
                _clear_context_keys(context)
                await CallbackHandlers._show_groups_list(
                    update, context, query, user_id, lang)
                return True

            if data.startswith("back_to_groups:") or data.startswith("sec_back:"):
                StateManager.clear(user_id)
                _clear_context_keys(context)
                await CallbackHandlers._show_groups_list(
                    update, context, query, user_id, lang)
                return True

            if data.startswith("set_warn_count:"):
                parts = data.split(":")
                if len(parts) != 3:
                    await safe_edit(query, await _trans('invalid_data', lang, "❌"),
                                    bot=context.bot)
                    return True
                chat_id = _coerce_int(parts[1])
                count = _coerce_int(parts[2])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                if count < 1 or count > 100:
                    await safe_edit(query,
                        await _trans('invalid_number', lang, "❌"), bot=context.bot)
                    return True
                await DB.update_security_settings(chat_id, max_warnings=count)
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return True

            if data == f"{CB.POST_CLEAR}_confirm":
                active = await DB.get_active_channel(user_id)
                if not active:
                    await safe_edit(query,
                        await _trans('no_active_channel', lang, "❌"),
                        bot=context.bot)
                    return True
                count = await DB.fetchval(
                    "SELECT COUNT(*) FROM posts WHERE channel_db_id=?",
                    (active,), default=0)
                text = (await _trans('clear_posts_confirm_title', lang, "⚠️") + "\n\n" +
                        _fmt(await _trans('clear_posts_confirm_body', lang, "{count}"),
                             count=count) + "\n" +
                        await _trans('clear_posts_confirm_note', lang, ""))
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('yes_clear_all', lang, "🧹"),
                        callback_data=CB.POST_CLEAR)],
                    [InlineKeyboardButton(
                        await _trans('cancel_btn', lang, "❌"),
                        callback_data=CB.POST_LIST)],
                ])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return True

            if data.startswith("post_del_confirm:"):
                post_id = _coerce_int(data.split(":")[-1])
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('yes_delete_post', lang, "🗑️"),
                        callback_data=f"{CB.POST_DEL}:{post_id}")],
                    [InlineKeyboardButton(
                        await _trans('cancel_btn', lang, "❌"),
                        callback_data=CB.POST_LIST)],
                ])
                text = _fmt(await _trans('delete_post_confirm', lang, "⚠️"),
                            post_id=post_id)
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return True

            if data.startswith("ch_del_confirm:"):
                ch_id = _coerce_int(data.split(":")[-1])
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('yes_delete_channel', lang, "🗑️"),
                        callback_data=f"{CB.CH_DEL}:{ch_id}")],
                    [InlineKeyboardButton(
                        await _trans('cancel_btn', lang, "❌"),
                        callback_data=CB.CH_LIST)],
                ])
                await safe_edit(query,
                    await _trans('delete_channel_confirm', lang, "⚠️"),
                    reply_markup=kb, parse_mode='HTML', bot=context.bot)
                return True

            if data.startswith("grp_del_confirm:"):
                chat_id = _coerce_int(data.split(":")[-1])
                if not await is_authorized_in_group(context.bot, chat_id, user_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                if not await _is_group_owner(user_id, chat_id):
                    await safe_edit(query,
                        await _trans('not_owner', lang, "❌"), bot=context.bot)
                    return True
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('yes_delete_group', lang, "🗑️"),
                        callback_data=f"grp_del:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('cancel_btn', lang, "❌"),
                        callback_data=CB.GROUPS)],
                ])
                await safe_edit(query,
                    await _trans('delete_group_confirm', lang, "⚠️"),
                    reply_markup=kb, parse_mode='HTML', bot=context.bot)
                return True

            if data.startswith("set_warn_penalty:"):
                parts = data.split(":")
                if len(parts) != 3:
                    return True
                _, penalty_type, chat_id_str = parts
                chat_id = _coerce_int(chat_id_str)
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                if penalty_type not in DB.VALID_PENALTY_TYPES:
                    await safe_edit(query,
                        await _trans('invalid_penalty_type', lang, "❌"),
                        bot=context.bot)
                    return True
                await DB.update_security_settings(chat_id, warn_penalty=penalty_type)
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return True

            if data.startswith("set_duration:"):
                parts = data.split(":")
                if len(parts) < 4:
                    return True
                penalty_type = parts[1]
                chat_id = _coerce_int(parts[2])
                duration = _coerce_int(parts[3])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                col_map = {
                    'mute': 'mute_default_duration',
                    'ban': 'ban_default_duration',
                    'restrict': 'restrict_default_duration',
                    'antiflood': 'antiflood_penalty_duration',
                    'night': 'night_mode_action_duration',
                    'warn_penalty': 'warn_penalty_duration',
                    'delete_penalty': 'delete_penalty_duration',
                    'violation': 'violation_penalty_duration',
                }
                col = col_map.get(penalty_type)
                if col is None:
                    await safe_edit(query,
                        await _trans('invalid_penalty_type', lang, "❌"),
                        bot=context.bot)
                    return True
                await DB.update_security_settings(chat_id, **{col: duration})
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return True

            if data.startswith("sec_set_del_penalty:"):
                parts = data.split(":")
                if len(parts) != 3:
                    return True
                _, penalty_type, chat_id_str = parts
                chat_id = _coerce_int(chat_id_str)
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                if penalty_type == "none":
                    await DB.update_security_settings(chat_id, delete_penalty="none")
                elif penalty_type in DB.VALID_PENALTY_TYPES:
                    await DB.update_security_settings(chat_id, delete_penalty=penalty_type)
                else:
                    await safe_edit(query,
                        await _trans('invalid_penalty_type', lang, "❌"),
                        bot=context.bot)
                    return True
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return True

            if data.startswith("sec_set_del_penalty_duration:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_penalty_durations(
                    update, context, query, chat_id, lang, 'delete_penalty')
                return True

            if data.startswith("sec_penalty_durations:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_all_penalty_durations_menu(
                    query, context, chat_id, lang)
                return True

            for prefix, action_type in (
                ("sec_set_mute_duration:", "mute"),
                ("sec_set_ban_duration:", "ban"),
                ("sec_set_restrict_duration:", "restrict"),
                ("sec_antiflood_duration:", "antiflood"),
                ("sec_night_duration:", "night"),
            ):
                if data.startswith(prefix):
                    parts = data.split(":")
                    if len(parts) != 2:
                        return True
                    chat_id = _coerce_int(parts[1])
                    if not await _check_sec_auth(context, user_id, chat_id):
                        await safe_edit(query,
                            await _trans('no_permission', lang, "❌"),
                            bot=context.bot)
                        return True
                    await CallbackHandlers._show_penalty_durations(
                        update, context, query, chat_id, lang, action_type)
                    return True

            if data.startswith("sec_warn_penalty_duration:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_penalty_durations(
                    update, context, query, chat_id, lang, 'warn_penalty')
                return True

            if data.startswith("sec_warn_penalty:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_warn_penalty_types(
                    update, context, query, chat_id, lang)
                return True

            if data.startswith("sec_warn_count:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_warn_count_buttons(
                    update, context, query, chat_id, lang)
                return True

            if data.startswith("sec_warn_toggle:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                settings = await CallbackHandlers._get_security_settings_cached(chat_id)
                new_val = 1 - _coerce_int(settings.get('warn_enabled', 0))
                await DB.update_security_settings(chat_id, warn_enabled=new_val)
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return True

            if data.startswith("sec_penalty_"):
                parts = data.split(":")
                if len(parts) < 2:
                    return True
                if parts[1].lstrip('-').isdigit():
                    chat_id = int(parts[1])
                else:
                    chat_id = await _resolve_sec_chat_id(context, data)
                if chat_id is None:
                    await safe_edit(query,
                        await _trans('group_not_specified', lang, "❌"),
                        bot=context.bot)
                    return True
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                action = parts[0][4:] if parts[0].startswith("sec_") else parts[0]
                action = action.replace("penalty_", "", 1)
                if action in ('ban', 'mute', 'kick', 'restrict', 'none'):
                    await DB.update_security_settings(chat_id, auto_penalty=action)
                    await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                    await CallbackHandlers._refresh_security_view(
                        query, context, chat_id, lang)
                else:
                    await safe_edit(query,
                        await _trans('invalid_penalty_type', lang, "❌"),
                        bot=context.bot)
                return True

            for prefix, state, prompt_key, prompt_default in (
                ("sec_set_antiflood_messages:", UserState.WAIT_ANTIFLOOD_MESSAGES,
                 "send_antiflood_messages", "📊 أرسل عدد الرسائل:"),
                ("sec_set_antiflood_seconds:", UserState.WAIT_ANTIFLOOD_SECONDS,
                 "send_antiflood_seconds", "⏱️ أرسل عدد الثواني:"),
            ):
                if data.startswith(prefix):
                    parts = data.split(":")
                    if len(parts) != 2:
                        return True
                    chat_id = _coerce_int(parts[1])
                    if not await _check_sec_auth(context, user_id, chat_id):
                        await safe_edit(query,
                            await _trans('no_permission', lang, "❌"),
                            bot=context.bot)
                        return True
                    StateManager.set(user_id, state)
                    _set_sec_chat(context, chat_id)
                    await safe_edit(query,
                        await _trans(prompt_key, lang, prompt_default),
                        bot=context.bot)
                    return True

            if data.startswith("sec_antiflood_penalty:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_penalty_type_selection(
                    update, context, query, chat_id, lang, 'antiflood_penalty')
                return True

            if data.startswith("sec_set_antiflood_penalty:"):
                parts = data.split(":")
                if len(parts) < 3:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                penalty_type = parts[2]
                if penalty_type in ('ban', 'mute', 'kick', 'restrict', 'none'):
                    await DB.update_security_settings(chat_id, antiflood_penalty=penalty_type)
                    await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                    await CallbackHandlers._refresh_security_view(
                        query, context, chat_id, lang)
                return True

            for prefix, state, prompt_key, prompt_default in (
                ("sec_set_night_start:", UserState.WAIT_NIGHT_START,
                 "send_night_start", "🌙 أرسل وقت البدء (HH:MM):"),
                ("sec_set_night_end:", UserState.WAIT_NIGHT_END,
                 "send_night_end", "🌙 أرسل وقت النهاية (HH:MM):"),
            ):
                if data.startswith(prefix):
                    parts = data.split(":")
                    if len(parts) != 2:
                        return True
                    chat_id = _coerce_int(parts[1])
                    if not await _check_sec_auth(context, user_id, chat_id):
                        await safe_edit(query,
                            await _trans('no_permission', lang, "❌"),
                            bot=context.bot)
                        return True
                    StateManager.set(user_id, state)
                    _set_sec_chat(context, chat_id)
                    await safe_edit(query,
                        await _trans(prompt_key, lang, prompt_default),
                        bot=context.bot)
                    return True

            if data.startswith("sec_night_action:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_penalty_type_selection(
                    update, context, query, chat_id, lang, 'night_action')
                return True

            if data.startswith("sec_set_night_action:"):
                parts = data.split(":")
                if len(parts) < 3:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                action_type = parts[2]
                if action_type in ('ban', 'mute', 'kick', 'restrict'):
                    await DB.update_security_settings(
                        chat_id, night_mode_action=action_type)
                    await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                    await CallbackHandlers._refresh_security_view(
                        query, context, chat_id, lang)
                return True

            if data.startswith("sec_violation_settings:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_violation_penalties(
                    update, context, query, chat_id, lang)
                return True

            if data.startswith("sec_set_violation_strikes:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                StateManager.set(user_id, UserState.WAIT_VIOLATION_STRIKES)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_violation_strikes', lang, "🔢"),
                    bot=context.bot)
                return True

            if data.startswith("sec_set_violation_duration:"):
                parts = data.split(":")
                if len(parts) != 2:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                await CallbackHandlers._show_penalty_durations(
                    update, context, query, chat_id, lang, 'violation')
                return True

            if data.startswith("sec_set_violation_penalty:"):
                parts = data.split(":")
                if len(parts) < 3:
                    return True
                chat_id = _coerce_int(parts[1])
                if not await _check_sec_auth(context, user_id, chat_id):
                    await safe_edit(query,
                        await _trans('no_permission', lang, "❌"), bot=context.bot)
                    return True
                penalty_type = parts[2]
                if penalty_type in ('ban', 'mute', 'kick', 'restrict', 'none'):
                    await DB.update_security_settings(
                        chat_id, violation_penalty=penalty_type)
                    await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                    await CallbackHandlers._refresh_security_view(
                        query, context, chat_id, lang)
                return True

            if data.startswith("buy_sub_"):
                await CallbackHandlers._handle_buy_subscription(
                    update, context, query, user_id, data, lang)
                return True

            if data.startswith("buy_gift:"):
                await CallbackHandlers._handle_buy_gift(
                    update, context, query, user_id, data, lang)
                return True

            if data.startswith("grp_del:"):
                await CallbackHandlers._handle_group_delete(
                    update, context, query, user_id, data, lang)
                return True

            if data.startswith(CB.GRP_SET + ":"):
                await CallbackHandlers._handle_group_settings(
                    update, context, query, user_id, lang, data)
                return True

            if data.startswith(CB.CH_SEL + ":"):
                await CallbackHandlers._handle_channel_select(
                    update, context, query, user_id, data, lang)
                return True

            if data.startswith(CB.CH_DEL + ":"):
                await CallbackHandlers._handle_channel_delete(
                    update, context, query, user_id, lang, data)
                return True

            if data.startswith(CB.CH_STATS + ":"):
                await CallbackHandlers._handle_channel_stats(
                    update, context, query, user_id, data, lang)
                return True

            if data.startswith(CB.POST_DEL + ":"):
                await CallbackHandlers._handle_post_delete(
                    update, context, query, user_id, lang, data)
                return True

            return False

        except BadRequest as e:
            if "query is too old" not in str(e).lower():
                logger.error(f"❌ param BadRequest: {e}", exc_info=True)
            return True
        except Exception as e:
            logger.error(f"❌ param: {e}", exc_info=True)
            try:
                await safe_edit(query,
                    await _trans('unexpected_error', lang, "❌"),
                    bot=context.bot)
            except Exception:
                pass
            return True

    # ═════════════════════════════════════════════════════════════
    # Refresh
    # =================================================================

    @staticmethod
    async def _refresh_security_view(query, context, chat_id, lang):
        try:
            await CallbackHandlers._invalidate_security_settings_cache(chat_id)
            await CallbackHandlers._render_security_two_phase(
                query, context, chat_id, lang, force_refresh_settings=False)
        except Exception as e:
            logger.error(f"_refresh_security_view: {e}", exc_info=True)

    # ═════════════════════════════════════════════════════════════
    # معالجات صغيرة
    # =================================================================

    @staticmethod
    async def _render_settings(query, context, user_id, lang):
        try:
            s = await DB.get_user_settings_batch(user_id) or {}
            if not isinstance(s, dict):
                s = _row_to_dict(s) or {}
            auto = "✅" if s.get('auto_publish') else "❌"
            rec = "✅" if s.get('auto_recycle') else "❌"
            kb = KeyboardFactory.build("settings", lang=lang)
            text = await _trans('settings_title', lang, "⚙️")
            auto_label = await _trans('auto_publish_status', lang, "📤")
            rec_label = await _trans('auto_recycle_status', lang, "♻️")
            text = f"{text}\n\n{auto_label}: {auto}\n{rec_label}: {rec}"
            await safe_edit(query, text, reply_markup=kb, bot=context.bot)
        except Exception as e:
            logger.error(f"_render_settings: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _render_referral(query, context, user_id, lang):
        try:
            stats = await DB.get_referral_stats(user_id) or {}
            if not isinstance(stats, dict):
                stats = _row_to_dict(stats) or {}
            code = await DB.get_referral_code(user_id)
            if code and code.startswith('ref_'):
                code = code[4:]
            link = f"https://t.me/{CONFIG.BOT_USERNAME}?start=ref_{code}"
            title = await _trans('referral_title', lang, "🔗")
            link_label = await _trans('referral_link_label', lang, "📎")
            referred_label = await _trans('referral_referred', lang, "👥")
            available_label = await _trans('referral_available', lang, "🎁")
            text = (
                f"{title}\n\n{link_label}\n{link}\n\n"
                f"{referred_label} {stats.get('total', 0)}\n"
                f"{available_label} {stats.get('available', 0)}"
            )
            back = KeyboardFactory.get_text("back", lang)
            claim = await _trans('ref_claim', lang, "🎁")
            list_l = await _trans('ref_list', lang, "📋")
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton(claim, callback_data=CB.REF_CLAIM),
                 InlineKeyboardButton(list_l, callback_data=CB.REF_LIST)],
                [InlineKeyboardButton(back, callback_data=CB.BACK)],
            ])
            await safe_edit(query, text, reply_markup=kb, bot=context.bot)
        except Exception as e:
            logger.error(f"_render_referral: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _render_reminder(query, context, user_id, lang):
        try:
            settings = await DB.get_reminder_settings(user_id) or {}
            if not isinstance(settings, dict):
                settings = _row_to_dict(settings) or {}
            title = await _trans('reminder_title', lang, "⏰")
            sub_l = await _trans('rem_sub_label', lang, "🔔")
            daily_l = await _trans('rem_daily_label', lang, "📊")
            weekly_l = await _trans('rem_weekly_label', lang, "📈")
            text = (
                f"{title}\n\n"
                f"{sub_l}: {'✅' if settings.get('subscription_reminder') else '❌'}\n"
                f"{daily_l}: {'✅' if settings.get('daily_stats_reminder') else '❌'}\n"
                f"{weekly_l}: {'✅' if settings.get('weekly_report') else '❌'}"
            )
            await safe_edit(query, text,
                            reply_markup=KeyboardFactory.build("reminder", lang=lang),
                            bot=context.bot)
        except Exception as e:
            logger.error(f"_render_reminder: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _handle_reminder_toggle(query, context, user_id, lang, base_data):
        try:
            settings = await DB.get_reminder_settings(user_id) or {}
            if not isinstance(settings, dict):
                settings = _row_to_dict(settings) or {}
            if base_data == CB.REM_TOGGLE_SUB:
                new_val = not settings.get('subscription_reminder', False)
                await DB.update_reminder_settings(user_id, subscription_reminder=new_val)
            elif base_data == CB.REM_TOGGLE_DAILY:
                new_val = not settings.get('daily_stats_reminder', False)
                await DB.update_reminder_settings(user_id, daily_stats_reminder=new_val)
            elif base_data == CB.REM_TOGGLE_WEEKLY:
                new_val = not settings.get('weekly_report', False)
                await DB.update_reminder_settings(user_id, weekly_report=new_val)
            await CallbackHandlers._render_reminder(query, context, user_id, lang)
        except Exception as e:
            logger.error(f"_handle_reminder_toggle: {e}", exc_info=True)

    # ═════════════════════════════════════════════════════════════
    # قائمة الترجمة
    # =================================================================

    @staticmethod
    async def _render_translation_menu(query, context, user_id, lang):
        try:
            current_lang = await DB.get_user_language(user_id) or 'ar'
        except Exception:
            current_lang = 'ar'
        try:
            langs = TranslationManager.get_available_languages() or {}
        except Exception:
            langs = {}

        is_off = (current_lang == 'off')
        title = await _trans('translation', lang, "🌐")

        if is_off:
            state_line = await _trans('translation_off', lang, "❌")
            hint_line = await _trans('choose_language', lang, "💡")
        else:
            current_name = langs.get(current_lang, current_lang)
            state_line = f"✅ {_html.escape(str(current_name))}"
            hint_line = await _trans('choose_language', lang, "💡")

        text = (f"{title}\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
                f"{state_line}\n<i>{hint_line}</i>")

        rows: list = []
        current_row: list = []
        for code, name in langs.items():
            label = f"✅ {name}" if (code == current_lang and not is_off) else name
            current_row.append((label, code))
            if len(current_row) == 2:
                rows.append(current_row)
                current_row = []
        if current_row:
            rows.append(current_row)

        kb = [[InlineKeyboardButton(t, callback_data=f"lang_{c}") for t, c in row]
              for row in rows]

        if not is_off:
            kb.append([InlineKeyboardButton(
                await _trans('translation_off', lang, "❌"),
                callback_data=CB.TRANS_OFF)])

        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang), callback_data=CB.BACK)])

        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(kb),
                        parse_mode='HTML', bot=context.bot)

    @staticmethod
    async def _handle_language_change(update, context, query, user_id):
        data = query.data or ""
        lang_set = data[5:] if data.startswith("lang_") else data.split("_")[-1]
        valid_langs = {
            'ar', 'en', 'fr', 'tr', 'zh', 'ru', 'de', 'es', 'it',
            'pt', 'ja', 'ko', 'fa', 'ur', 'nl', 'pl', 'hi', 'off',
        }
        if lang_set in valid_langs:
            await DB.set_user_language(user_id, lang_set)
            await invalidate_user_cache(user_id)
            try:
                context.user_data.pop('lang', None)
            except Exception:
                pass
            try:
                from handlers_message import clear_lang_cache
                clear_lang_cache(context)
            except (ImportError, AttributeError) as e:
                logger.warning(f"clear_lang_cache fallback: {e}")
                try:
                    context.user_data.pop('translation_cache', None)
                except Exception:
                    pass
            ok = await CallbackHandlers._show_main_menu_inline(query, context, user_id)
            if not ok:
                await CommandHandlers.start(update, context)
        else:
            lang = await DB.get_user_language(user_id) or 'ar'
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # معالجات فرعية
    # =================================================================

    @staticmethod
    async def _handle_buy_subscription(update, context, query, user_id, data, lang):
        try:
            days = int(data.split("_")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        plan_names = {1: "يوم", 7: "أسبوع", 30: "شهر", 90: "3 أشهر", 365: "سنة"}
        plan_name = plan_names.get(days)
        if not plan_name:
            await safe_edit(query,
                await _trans('plan_not_found', lang, "❌"), bot=context.bot)
            return
        plan = await DB.get_plan_by_name(plan_name)
        plan_d = _row_to_dict(plan)
        if not plan_d:
            try:
                fallback = await DB.fetchone(
                    "SELECT * FROM plans WHERE duration_days = ? "
                    "AND is_active = 1 AND is_gift = 0 LIMIT 1", (days,))
                plan_d = _row_to_dict(fallback)
            except Exception:
                pass
        if not plan_d:
            await safe_edit(query,
                await _trans('plan_not_found', lang, "❌"), bot=context.bot)
            return
        plan_id = plan_d.get('id', 0)
        price = _coerce_int(plan_d.get('price'), 0)
        name = plan_d.get('name') or plan_name
        description = plan_d.get('description') or name
        invoice_number = await DB.create_invoice(user_id, plan_id, price)
        if not invoice_number:
            await safe_edit(query,
                await _trans('invoice_failed', lang, "❌"), bot=context.bot)
            return
        try:
            await context.bot.send_invoice(
                chat_id=user_id, title=f"💎 {name}",
                description=description,
                payload=json.dumps({
                    'plan_id': plan_id, 'invoice': invoice_number,
                    'type': 'subscription'}),
                provider_token="", currency="XTR",
                prices=[LabeledPrice(name, price)],
            )
            await safe_delete_message(query)
        except Exception as e:
            logger.error(f"❌ invoice: {e}")
            try:
                await DB.execute(
                    "UPDATE invoices SET status='cancelled' WHERE number=?",
                    (invoice_number,))
            except Exception:
                pass
            await safe_edit(query, f"❌ {str(e)[:50]}", bot=context.bot)

    @staticmethod
    async def _handle_buy_gift(update, context, query, user_id, data, lang):
        try:
            gift_plan_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        plan = await DB.get_gift_plan(gift_plan_id)
        plan_d = _row_to_dict(plan)
        if not plan_d:
            await safe_edit(query,
                await _trans('plan_not_found', lang, "❌"), bot=context.bot)
            return
        plan_id = plan_d.get('id', 0)
        price = _coerce_int(plan_d.get('price'), 0)
        name = plan_d.get('name') or 'Gift'
        description = plan_d.get('description') or "Gift code"
        invoice_number = await DB.create_invoice(user_id, plan_id, price)
        if not invoice_number:
            await safe_edit(query,
                await _trans('invoice_failed', lang, "❌"), bot=context.bot)
            return
        try:
            await context.bot.send_invoice(
                chat_id=user_id, title=f"🎁 {name}",
                description=description,
                payload=json.dumps({
                    'gift_plan_id': plan_id, 'invoice': invoice_number,
                    'type': 'gift'}),
                provider_token="", currency="XTR",
                prices=[LabeledPrice(name, price)],
            )
            await safe_delete_message(query)
        except Exception as e:
            logger.error(f"❌ gift invoice: {e}")
            try:
                await DB.execute(
                    "UPDATE invoices SET status='cancelled' WHERE number=?",
                    (invoice_number,))
            except Exception:
                pass
            await safe_edit(query, f"❌ {str(e)[:50]}", bot=context.bot)

    @staticmethod
    async def _handle_group_delete(update, context, query, user_id, data, lang):
        try:
            chat_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        if not await is_authorized_in_group(context.bot, chat_id, user_id):
            await safe_edit(query,
                await _trans('no_permission', lang, "❌"), bot=context.bot)
            return
        if not await _is_group_owner(user_id, chat_id):
            await safe_edit(query,
                await _trans('not_owner', lang, "❌"), bot=context.bot)
            return
        try:
            await _invalidate_log_channel_menu_cache(chat_id)
        except Exception:
            pass
        if await DB.delete_group(chat_id):
            await _invalidate_after_channel_change(user_id)
            await safe_edit(query,
                await _trans('group_deleted', lang, "✅"), bot=context.bot)
        else:
            await safe_edit(query,
                await _trans('delete_failed', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _handle_group_settings(update, context, query, user_id, lang, data):
        try:
            chat_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        if not await _check_sec_auth(context, user_id, chat_id):
            await safe_edit(query,
                await _trans('no_permission', lang, "❌"), bot=context.bot)
            return
        _set_sec_chat(context, chat_id)
        await CallbackHandlers._render_security_two_phase(
            query, context, chat_id, lang, force_refresh_settings=False)

    @staticmethod
    async def _handle_channel_select(update, context, query, user_id, data, lang):
        try:
            ch_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        if await DB.set_active_channel(user_id, ch_id):
            await _invalidate_after_channel_change(user_id, ch_id)
            await safe_edit(query,
                await _trans('channel_selected', lang, "✅"), bot=context.bot)
        else:
            await safe_edit(query,
                await _trans('channel_select_failed', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _handle_channel_delete(update, context, query, user_id, lang, data):
        try:
            ch_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        if await DB.delete_channel(user_id, ch_id):
            context.user_data['channel_page'] = 0
            await _invalidate_after_channel_change(user_id, ch_id)
            await CallbackHandlers._show_channel_list(
                update, context, query, user_id, lang)
        else:
            await safe_edit(query,
                await _trans('delete_failed', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _handle_channel_stats(update, context, query, user_id, data, lang):
        try:
            ch_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        try:
            stats = await DB.get_channel_stats(user_id, ch_id) or {}
            if not isinstance(stats, dict):
                stats = _row_to_dict(stats) or {}
            text = (f"{await _trans('channel_stats', lang, '📊')}\n\n"
                    f"📝 {stats.get('total', 0)}\n"
                    f"✅ {stats.get('published', 0)}\n"
                    f"⏳ {stats.get('unpublished', 0)}")
        except Exception as e:
            logger.error(f"_handle_channel_stats: {e}")
            text = await _trans('error_occurred', lang, "❌")
        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
            [[InlineKeyboardButton(KeyboardFactory.get_text("back", lang),
                                   callback_data=CB.CH_LIST)]]),
            bot=context.bot)

    @staticmethod
    async def _handle_post_add(update, context, query, user_id):
        lang = 'ar'
        try:
            lang = await DB.get_user_language(user_id) or 'ar'
        except Exception:
            pass
        try:
            is_owner = _is_primary_owner(user_id)
            has_sub = True if is_owner else await DB.has_active_subscription(user_id)
            active = await DB.get_active_channel(user_id)
            if not has_sub:
                await safe_edit(query,
                    await _trans('subscription_expired', lang, "❌"), bot=context.bot)
                return
            if not active:
                await safe_edit(query,
                    await _trans('no_active_channel_alert', lang, "❌"),
                    parse_mode='HTML', bot=context.bot)
                return
            StateManager.set(user_id, UserState.ADDING_POSTS)
            finish_btn = await _trans('finish_btn', lang, "✅")
            await safe_edit(
                query, await _trans('send_posts_prompt', lang, "📥"),
                reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(finish_btn, callback_data="finish_posts")]]),
                bot=context.bot)
        except Exception as e:
            logger.error(f"❌ _handle_post_add: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _handle_post_publish(update, context, query, user_id):
        lang = 'ar'
        try:
            lang = await DB.get_user_language(user_id) or 'ar'
        except Exception:
            pass
        is_owner = _is_primary_owner(user_id)
        if not is_owner and not await DB.has_active_subscription(user_id):
            await safe_edit(query,
                await _trans('subscription_expired', lang, "❌"), bot=context.bot)
            return
        active = await DB.get_active_channel(user_id)
        if not active:
            await safe_edit(query,
                await _trans('no_channels', lang, "❌"), bot=context.bot)
            return
        raw_result = await DB.get_next_post(active)
        post, was_recycled = CallbackHandlers._unwrap_get_next_post(raw_result)
        if not post or not isinstance(post, dict):
            await safe_edit(query,
                await _trans('no_posts', lang, "📭"), bot=context.bot)
            return
        ch_info = await DB.get_channel_info(user_id, active)
        ch_info_d = _row_to_dict(ch_info)
        if not ch_info_d or not ch_info_d.get('channel_id'):
            await safe_edit(query,
                await _trans('channel_info_unavailable', lang, "❌"), bot=context.bot)
            return
        bot = context.bot
        ch_id = ch_info_d['channel_id']
        recycled_flag = was_recycled

        async def _publish_task():
            try:
                async with _publish_semaphore:
                    result = await CallbackHandlers._publish_single(
                        bot, active, ch_id, post)
                if result:
                    try:
                        await _invalidate_after_channel_change(user_id, active)
                    except Exception:
                        pass
                    if recycled_flag:
                        msg = await _trans('publish_success_recycled', lang, "✅")
                    else:
                        msg = await _trans('publish_success', lang, "✅")
                    await safe_send(bot, user_id, msg)
                else:
                    if recycled_flag:
                        msg = await _trans('publish_failed_recycled', lang, "❌")
                    else:
                        msg = await _trans('publish_failed', lang, "❌")
                    await safe_send(bot, user_id, msg)
            except Exception as e:
                logger.error(f"❌ _publish_task: {e}", exc_info=True)
                try:
                    msg = _fmt(await _trans('publish_unexpected_error', lang, "❌"),
                               error=str(e)[:80])
                    await safe_send(bot, user_id, msg)
                except Exception:
                    pass

        task = asyncio.create_task(_publish_task())
        ACTIVE_TASKS.add(task)
        task.add_done_callback(ACTIVE_TASKS.discard)
        if was_recycled:
            msg = await _trans('publish_started_recycled', lang, "✅")
        else:
            msg = await _trans('publish_started', lang, "✅")
        await safe_edit(query, msg, bot=context.bot)

    @staticmethod
    async def _handle_post_delete(update, context, query, user_id, lang, data):
        try:
            post_id = int(data.split(":")[-1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        active = await DB.get_active_channel(user_id)
        if active and await DB.delete_post(user_id, post_id, active):
            await _invalidate_after_channel_change(user_id, active)
            await CallbackHandlers._show_post_list(
                update, context, query, user_id, lang)
        else:
            await safe_edit(query,
                await _trans('delete_failed', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _handle_publish_all(update, context, query, user_id):
        lang = 'ar'
        try:
            lang = await DB.get_user_language(user_id) or 'ar'
        except Exception:
            pass
        is_owner = _is_primary_owner(user_id)
        if not is_owner and not await DB.has_active_subscription(user_id):
            await safe_edit(query,
                await _trans('subscription_expired', lang, "❌"), bot=context.bot)
            return
        channels = await DB.get_user_channels(user_id)
        if not channels:
            await safe_edit(query,
                await _trans('no_channels_for_publish', lang, "❌"), bot=context.bot)
            return
        task = asyncio.create_task(
            CallbackHandlers._publish_all(context.bot, user_id, channels))
        ACTIVE_TASKS.add(task)
        task.add_done_callback(ACTIVE_TASKS.discard)
        await safe_edit(query,
            await _trans('publish_all_started', lang, "✅"), bot=context.bot)

    @staticmethod
    async def _show_groups_list(update, context, query, user_id, lang):
        groups = await DB.get_user_groups(user_id)
        if not groups:
            back = KeyboardFactory.get_text("back", lang)
            add_bot = await _trans('add_bot_to_group', lang, "➕")
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton(
                    add_bot,
                    url=f"https://t.me/{CONFIG.BOT_USERNAME}?startgroup")],
                [InlineKeyboardButton(back, callback_data=CB.BACK)],
            ])
            await safe_edit(query,
                await _trans('no_groups', lang, "📭"),
                reply_markup=kb, bot=context.bot)
            return

        title = await _trans('groups_title', lang, "👥")
        text = f"{title}\n\n"
        kb = []
        display_idx = 0
        first_chat_id = None
        for g in groups:
            gd = _row_to_dict(g) or {}
            gid = gd.get('chat_id')
            if gid is None:
                continue
            display_idx += 1
            if display_idx == 1:
                first_chat_id = gid
            name = gd.get('chat_name') or f"Group {gid}"
            status = "⛔" if gd.get('banned') else "✅"
            number = _group_number(display_idx)
            text += f"{status} {number} {name}\n"
            sec_label = _fmt(await _trans('group_security_btn', lang, "{number} ⚙️"),
                             number=number)
            del_label = _fmt(await _trans('group_delete_btn', lang, "{number} 🗑️"),
                             number=number)
            kb.append([
                InlineKeyboardButton(sec_label,
                                     callback_data=f"{CB.GRP_SET}:{gid}"),
                InlineKeyboardButton(del_label,
                                     callback_data=f"grp_del_confirm:{gid}"),
            ])
        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang), callback_data=CB.BACK)])
        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(kb),
                        bot=context.bot)
        if first_chat_id:
            task = asyncio.create_task(
                CallbackHandlers._preload_group_security(first_chat_id))
            ACTIVE_TASKS.add(task)
            task.add_done_callback(ACTIVE_TASKS.discard)

    @staticmethod
    async def _preload_group_security(chat_id: int) -> None:
        try:
            key = f"sec_set_{chat_id}"
            cached = await _security_settings_cache.get(key)
            if cached is None:
                settings = await DB.get_security_settings(chat_id) or {}
                if not isinstance(settings, dict):
                    settings = _row_to_dict(settings) or {}
                await _security_settings_cache.set(key, settings,
                                                   ttl=SEC_SETTINGS_CACHE_TTL)
            stats_key = f"sec_stats_{chat_id}"
            cached_stats = await _security_stats_cache_local.get(stats_key)
            if cached_stats is None:
                stats = await KeyboardFactory._get_security_stats(chat_id) or {}
                await _security_stats_cache_local.set(stats_key, stats,
                                                      ttl=SEC_STATS_CACHE_TTL)
        except Exception as e:
            logger.debug(f"_preload_group_security({chat_id}): {e}")

    @staticmethod
    def _unwrap_get_next_post(result) -> Tuple[Optional[Dict], bool]:
        if result is None:
            return None, False
        if isinstance(result, tuple) and len(result) == 2:
            post_obj, was_recycled = result
            if post_obj is None:
                return None, bool(was_recycled)
            as_dict = _row_to_dict(post_obj)
            if as_dict is not None:
                return as_dict, bool(was_recycled)
            return None, False
        as_dict = _row_to_dict(result)
        if as_dict is not None:
            return as_dict, False
        return None, False

    @staticmethod
    async def _publish_single(bot, ch_db_id, ch_tele, post) -> bool:
        if isinstance(post, tuple) and len(post) == 2:
            post, _ = CallbackHandlers._unwrap_get_next_post(post)
        if not isinstance(post, dict):
            post = _row_to_dict(post)
        if not isinstance(post, dict):
            return False
        post_id = post.get('id')
        try:
            text = post.get('text', '') or ''
            media_type = post.get('media_type')
            media_file_id = post.get('media_file_id')
            if not text and not media_type and not media_file_id:
                return False
            caption = text[:MAX_CAPTION_LENGTH] if text else None
            if media_type == 'photo' and media_file_id:
                await bot.send_photo(ch_tele, media_file_id, caption=caption)
            elif media_type == 'video' and media_file_id:
                await bot.send_video(ch_tele, media_file_id, caption=caption)
            elif media_type == 'document' and media_file_id:
                await bot.send_document(ch_tele, media_file_id, caption=caption)
            elif media_type == 'audio' and media_file_id:
                await bot.send_audio(ch_tele, media_file_id, caption=caption)
            elif media_type == 'voice' and media_file_id:
                await bot.send_voice(ch_tele, media_file_id)
                if text:
                    try:
                        await bot.send_message(ch_tele, text)
                    except Exception:
                        pass
            elif media_type == 'animation' and media_file_id:
                await bot.send_animation(ch_tele, media_file_id, caption=caption)
            elif media_type == 'sticker' and media_file_id:
                await bot.send_sticker(ch_tele, media_file_id)
                if text:
                    try:
                        await bot.send_message(ch_tele, text)
                    except Exception:
                        pass
            elif media_type == 'video_note' and media_file_id:
                await bot.send_video_note(ch_tele, media_file_id)
                if text:
                    try:
                        await bot.send_message(ch_tele, text)
                    except Exception:
                        pass
            else:
                if text and len(text) > MAX_MESSAGE_LENGTH:
                    for i in range(0, len(text), MAX_MESSAGE_LENGTH):
                        await bot.send_message(ch_tele, text[i:i + MAX_MESSAGE_LENGTH])
                else:
                    await bot.send_message(ch_tele, text if text else ".")
            if post_id:
                await DB.mark_post_published(post_id)
            await DB.update_last_publish(ch_db_id)
            await DB.update_next_publish(ch_db_id)
            return True
        except RetryAfter as e:
            delay = e.retry_after
            if isinstance(delay, timedelta):
                delay = delay.total_seconds()
            try:
                await asyncio.sleep(float(delay))
            except Exception:
                await asyncio.sleep(5)
            if post_id:
                try:
                    await DB.increment_post_fail(post_id)
                except Exception:
                    pass
            return False
        except Forbidden:
            try:
                await DB.execute(
                    "UPDATE user_channels SET banned=1 WHERE id=?", (ch_db_id,))
            except Exception:
                pass
            if post_id:
                try:
                    await DB.increment_post_fail(post_id)
                except Exception:
                    pass
            return False
        except Exception as e:
            logger.error(f"❌ publish: {e}", exc_info=True)
            if post_id:
                try:
                    await DB.increment_post_fail(post_id)
                except Exception:
                    pass
            return False

    @staticmethod
    async def _publish_all(bot, user_id, channels):
        lang = 'ar'
        try:
            lang = await DB.get_user_language(user_id) or 'ar'
        except Exception:
            pass
        published = 0
        failed = 0
        tasks = []
        banned_count = 0
        no_post_count = 0
        lost_count = 0
        try:
            for ch in channels:
                chd = _row_to_dict(ch) or {}
                if chd.get('banned'):
                    banned_count += 1
                    continue
                raw_result = await DB.get_next_post(chd.get('id'))
                post, _ = CallbackHandlers._unwrap_get_next_post(raw_result)
                if post and isinstance(post, dict):
                    ch_info = await DB.get_channel_info(user_id, chd.get('id'))
                    ch_info_d = _row_to_dict(ch_info)
                    if ch_info_d and ch_info_d.get('channel_id'):
                        tasks.append((chd.get('id'), ch_info_d['channel_id'], post))
                    else:
                        pid = post.get('id')
                        if pid:
                            try:
                                await DB.increment_post_fail(pid)
                            except Exception:
                                pass
                        lost_count += 1
                else:
                    no_post_count += 1

            if not tasks:
                if banned_count == len(channels):
                    msg = await _trans('all_channels_banned', lang, "❌")
                elif no_post_count == len(channels) - banned_count:
                    msg = await _trans('no_valid_posts_short', lang, "📭")
                else:
                    msg = await _trans('no_valid_posts', lang, "📭")
                await safe_send(bot, user_id, msg)
                return

            async def run(task):
                async with _publish_semaphore:
                    try:
                        await asyncio.wait_for(
                            PUBLISH_RATE_LIMITER.acquire(),
                            timeout=PUBLISH_ACQUIRE_TIMEOUT)
                    except Exception:
                        pass
                    result = await CallbackHandlers._publish_single(
                        bot, task[0], task[1], task[2])
                    await asyncio.sleep(CallbackHandlers.PUBLISH_DELAY_SECONDS)
                    return result

            BATCH = CallbackHandlers.PUBLISH_BATCH_SIZE
            for i in range(0, len(tasks), BATCH):
                batch = tasks[i:i + BATCH]
                results = await asyncio.gather(
                    *(run(t) for t in batch), return_exceptions=True)
                for r in results:
                    if r is True:
                        published += 1
                    else:
                        failed += 1

            if lost_count:
                summary = _fmt(await _trans('publish_summary_lost', lang,
                                            "✅ {published} ❌ {failed} ⚠️ {lost}"),
                               published=published, failed=failed, lost=lost_count)
            else:
                summary = _fmt(await _trans('publish_summary', lang,
                                            "✅ {published} ❌ {failed}"),
                               published=published, failed=failed)
            await safe_send(bot, user_id, summary)

            try:
                await _invalidate_after_channel_change(
                    user_id, channel_db_id=None, invalidate_posts=False)
            except Exception:
                pass
        except Exception as e:
            logger.error(f"❌ _publish_all: {e}", exc_info=True)
            try:
                msg = _fmt(await _trans('publish_all_failed', lang, "❌"),
                           error=str(e)[:100])
                await safe_send(bot, user_id, msg)
            except Exception:
                pass

    @staticmethod
    async def _show_channel_list(update, context, query, user_id, lang=None):
        if not lang:
            lang = await DB.get_user_language(user_id) or 'ar'
        channels = await DB.get_user_channels(user_id)
        if not channels:
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton(
                    KeyboardFactory.get_text("ch_add", lang),
                    callback_data=CB.CH_ADD)],
                [InlineKeyboardButton(
                    KeyboardFactory.get_text("back", lang),
                    callback_data=CB.BACK)],
            ])
            await safe_edit(query,
                await _trans('no_channels', lang, "📭"),
                reply_markup=kb, bot=context.bot)
            return

        page = _coerce_int(context.user_data.get('channel_page'), 0)
        per_page = 5
        total_pages = max(1, (len(channels) + per_page - 1) // per_page)
        if page >= total_pages:
            page = total_pages - 1
        if page < 0:
            page = 0
        context.user_data['channel_page'] = page

        page_channels = channels[page * per_page:(page + 1) * per_page]
        start_index = page * per_page

        title = await _trans('your_channels', lang, "📡")
        text = f"{title} ({len(channels)})\n<i>{page + 1}/{total_pages}</i>\n\n"

        kb = []
        display_idx = start_index
        for ch in page_channels:
            chd = _row_to_dict(ch) or {}
            ch_id = chd.get('id')
            if ch_id is None:
                continue
            display_idx += 1
            number = _group_number(display_idx)
            st = "✅" if not chd.get('banned') else "🚫"
            raw_name = chd.get('channel_name') or f"Channel {display_idx}"
            name = _html.escape(str(raw_name)[:35])
            text += f"{st} {number}  {name}\n"
            kb.append([
                InlineKeyboardButton(f"{number} 📌",
                                     callback_data=f"{CB.CH_SEL}:{ch_id}"),
                InlineKeyboardButton(f"{number} 📅",
                                     callback_data=f"sched_open:{ch_id}"),
            ])
            kb.append([
                InlineKeyboardButton(f"{number} 📊",
                                     callback_data=f"{CB.CH_STATS}:{ch_id}"),
                InlineKeyboardButton(f"{number} 🗑️",
                                     callback_data=f"ch_del_confirm:{ch_id}"),
            ])

        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton(
                await _trans('prev_btn', lang, "⬅️"),
                callback_data="ch_page_prev"))
        if page < total_pages - 1:
            nav.append(InlineKeyboardButton(
                await _trans('next_btn', lang, "➡️"),
                callback_data="ch_page_next"))
        if nav:
            kb.append(nav)

        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("ch_add", lang), callback_data=CB.CH_ADD)])
        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang), callback_data=CB.BACK)])

        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(kb),
                        parse_mode='HTML', bot=context.bot)

    @staticmethod
    async def _show_post_list(update, context, query, user_id, lang=None):
        if not lang:
            lang = await DB.get_user_language(user_id) or 'ar'
        active = await DB.get_active_channel(user_id)
        if not active:
            await safe_edit(query,
                await _trans('no_active_channel', lang, "❌"), bot=context.bot)
            return
        per_page = 5
        total = await DB.fetchval(
            "SELECT COUNT(*) FROM posts WHERE channel_db_id=?",
            (active,), default=0)
        total_pages = max(1, (total + per_page - 1) // per_page)
        page = _coerce_int(context.user_data.get('post_page'), 0)
        if page >= total_pages:
            page = total_pages - 1
        if page < 0:
            page = 0
        context.user_data['post_page'] = page
        posts = await DB.fetchall(
            "SELECT id, text, published FROM posts WHERE channel_db_id=? "
            "ORDER BY created_at ASC LIMIT ? OFFSET ?",
            (active, per_page, page * per_page))
        title = await _trans('post_list', lang, "📋")
        text = f"{title} ({page + 1}/{total_pages})\n\n"
        kb = []
        for p in posts:
            pd = _row_to_dict(p) or {}
            pid = pd.get('id')
            if pid is None:
                continue
            text += f"🆔 {pid}: {(pd.get('text') or '')[:30]}\n"
            kb.append([InlineKeyboardButton(
                f"🗑️ {pid}",
                callback_data=f"post_del_confirm:{pid}")])
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton(
                await _trans('prev_btn', lang, "⬅️"),
                callback_data="post_page_prev"))
        if page < total_pages - 1:
            nav.append(InlineKeyboardButton(
                await _trans('next_btn', lang, "➡️"),
                callback_data="post_page_next"))
        if nav:
            kb.append(nav)
        kb.append([InlineKeyboardButton(
            await _trans('post_rec', lang, "🔄"), callback_data=CB.POST_REC)])
        kb.append([InlineKeyboardButton(
            await _trans('post_clear', lang, "🧹"),
            callback_data=f"{CB.POST_CLEAR}_confirm")])
        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang), callback_data=CB.BACK)])
        display_text = text if posts else await _trans('no_posts', lang, "📭")
        await safe_edit(query, display_text,
                        reply_markup=InlineKeyboardMarkup(kb), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Security handlers
    # ═════════════════════════════════════════════════════════════

    @staticmethod
    async def _handle_security(update, context, query, user_id, lang=None):
        if not lang:
            lang = await DB.get_user_language(user_id) or 'ar'
        data = query.data
        parts = data.split(":")

        chat_id = None
        if len(parts) >= 2 and parts[1].lstrip('-').isdigit():
            chat_id = int(parts[1])
        else:
            stored = (context.user_data.get('security_chat_id')
                      or context.user_data.get('sec_chat'))
            if stored:
                try:
                    chat_id = int(stored)
                except (TypeError, ValueError):
                    chat_id = None

        if chat_id is None:
            await safe_edit(query,
                await _trans('group_not_specified', lang, "❌"), bot=context.bot)
            return

        prefix0 = parts[0]
        action = prefix0[4:] if prefix0.startswith("sec_") else prefix0

        if not await _check_sec_auth(context, user_id, chat_id):
            await safe_edit(query,
                await _trans('no_permission', lang, "❌"), bot=context.bot)
            return

        try:
            if action == "log_channel_btn":
                await CallbackHandlers._show_log_channel_menu(
                    query, context, chat_id, user_id, lang)
                return

            if action == "auto_reply_menu":
                context.user_data['auto_chat'] = chat_id
                try:
                    kb = KeyboardFactory.build("auto_reply",
                                                chat_id=chat_id, lang=lang)
                except Exception:
                    kb = InlineKeyboardMarkup([
                        [InlineKeyboardButton(
                            KeyboardFactory.get_text("back", lang),
                            callback_data=f"{CB.GRP_SET}:{chat_id}")]])
                await safe_edit(query,
                    await _trans('auto_reply_title', lang, "🤖"),
                    reply_markup=kb, bot=context.bot)
                return

            if action == "maxlen":
                StateManager.set(user_id, UserState.WAIT_MAX_LEN)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_max_length', lang, "📏"), bot=context.bot)
                return

            if action == "act_log":
                await CallbackHandlers._show_admin_logs(
                    update, context, query, chat_id, lang)
                return

            if action in ("activate_all", "enable_all",
                          "deactivate_all", "disable_all"):
                is_activate = action in ("activate_all", "enable_all")
                confirm_action = ("activate_all_confirm" if is_activate
                                  else "deactivate_all_confirm")
                if is_activate:
                    confirm_text = await _trans('activate_all_confirmation',
                                                lang, "⚠️")
                else:
                    confirm_text = await _trans('deactivate_all_confirmation',
                                                lang, "⚠️")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('yes_btn', lang, "✅"),
                        callback_data=f"sec_{confirm_action}:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('cancel_btn', lang, "❌"),
                        callback_data=f"{CB.GRP_SET}:{chat_id}")],
                ])
                await safe_edit(query, confirm_text,
                                reply_markup=kb, bot=context.bot)
                return

            if action in ("activate_all_confirm", "deactivate_all_confirm"):
                is_activate = (action == "activate_all_confirm")
                activate_values = dict(
                    delete_links=1, delete_mentions=1, slow_mode=1,
                    slow_mode_seconds=5,
                    delete_videos=1, delete_audio=1, delete_animation=1,
                    delete_service=1, delete_documents=1, delete_stickers=1,
                    delete_forwarded=1, delete_polls=1, delete_games=1,
                    delete_voice=1, delete_video_note=1,
                    welcome_enabled=1, goodbye_enabled=1,
                    antiflood_enabled=1, antiflood_messages=5,
                    antiflood_seconds=5,
                    antiflood_penalty="mute", antiflood_penalty_duration=60,
                    night_mode_enabled=1,
                    night_mode_start="22:00", night_mode_end="06:00",
                    night_mode_action="mute", night_mode_action_duration=3600,
                    auto_approve_join=0, auto_reject_join=0, nsfw_enabled=0,
                    warn_enabled=1, max_warnings=3,
                    warn_penalty="mute", warn_penalty_duration=3600,
                    delete_banned_words=1, auto_penalty="mute",
                    delete_penalty="mute", delete_penalty_duration=3600,
                    violation_strikes=3, violation_duration=60,
                )
                deactivate_values = dict(
                    delete_links=0, delete_mentions=0, slow_mode=0,
                    slow_mode_seconds=0,
                    delete_videos=0, delete_audio=0, delete_animation=0,
                    delete_service=0, delete_documents=0, delete_stickers=0,
                    delete_forwarded=0, delete_polls=0, delete_games=0,
                    delete_voice=0, delete_video_note=0,
                    welcome_enabled=0, goodbye_enabled=0,
                    antiflood_enabled=0, antiflood_messages=0,
                    antiflood_seconds=0,
                    antiflood_penalty="none", antiflood_penalty_duration=0,
                    night_mode_enabled=0, night_mode_start="",
                    night_mode_end="",
                    night_mode_action="none", night_mode_action_duration=0,
                    auto_approve_join=0, auto_reject_join=0, nsfw_enabled=0,
                    warn_enabled=0, max_warnings=0,
                    warn_penalty="none", warn_penalty_duration=0,
                    delete_banned_words=0, auto_penalty="none",
                    delete_penalty="none", delete_penalty_duration=0,
                    violation_strikes=0, violation_duration=0,
                )
                values = activate_values if is_activate else deactivate_values
                action_name = (f"{'activate' if is_activate else 'deactivate'}_all_security")
                try:
                    async with DB.transaction() as _conn:
                        await DB.update_security_settings(chat_id, conn=_conn, **values)
                        await DB.add_admin_log(
                            chat_id=chat_id, admin_id=user_id,
                            action=action_name, target_id=None,
                            reason="", conn=_conn)
                except TypeError:
                    try:
                        await DB.update_security_settings(chat_id, **values)
                    except Exception as ex:
                        logger.error(f"activate/deactivate: {ex}", exc_info=True)
                        await safe_edit(query,
                            await _trans('error_occurred', lang, "❌"),
                            bot=context.bot)
                        return

                    async def _log_security_action():
                        try:
                            await DB.add_admin_log(
                                chat_id=chat_id, admin_id=user_id,
                                action=action_name)
                        except Exception:
                            pass
                    try:
                        log_task = asyncio.create_task(_log_security_action())
                        ACTIVE_TASKS.add(log_task)
                        log_task.add_done_callback(ACTIVE_TASKS.discard)
                    except Exception:
                        pass
                except Exception as ex:
                    logger.error(f"activate/deactivate: {ex}", exc_info=True)
                    await safe_edit(query,
                        await _trans('error_occurred', lang, "❌"), bot=context.bot)
                    return

                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return

            toggle_map = {
                "links": "delete_links", "mentions": "mentions",
                "slow": "slow_mode",
                "video": "delete_videos", "audio": "delete_audio",
                "anim": "delete_animation", "service": "delete_service",
                "doc": "delete_documents", "sticker": "delete_stickers",
                "forward": "delete_forwarded", "poll": "delete_polls",
                "game": "delete_games", "voice": "delete_voice",
                "videonote": "delete_video_note",
                "welcome": "welcome_enabled",
                "goodbye": "goodbye_enabled", "flood": "antiflood_enabled",
                "night": "night_mode_enabled",
                "approve_join": "auto_approve_join",
                "reject_join": "auto_reject_join", "nsfw": "nsfw_enabled",
            }

            if action in toggle_map:
                col = toggle_map[action]
                settings = await CallbackHandlers._get_security_settings_cached(chat_id)
                new_val = 1 - _coerce_int(settings.get(col, 0))
                update_data = {col: new_val}
                if action == "approve_join" and new_val:
                    update_data['auto_reject_join'] = 0
                elif action == "reject_join" and new_val:
                    update_data['auto_approve_join'] = 0
                await DB.update_security_settings(chat_id, **update_data)
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return

            if action == "warn":
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('warn_toggle_btn', lang, "✅"),
                        callback_data=f"sec_warn_toggle:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('warn_count_btn', lang, "🔢"),
                        callback_data=f"sec_warn_count:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('warn_penalty_btn', lang, "⚖️"),
                        callback_data=f"sec_warn_penalty:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('warn_penalty_duration_btn', lang, "⏱️"),
                        callback_data=f"sec_warn_penalty_duration:{chat_id}")],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=f"{CB.GRP_SET}:{chat_id}")],
                ])
                await safe_edit(query,
                    await _trans('warnings_management', lang, "⚠️"),
                    reply_markup=kb, bot=context.bot)
                return

            if action == "warn_penalty":
                await CallbackHandlers._show_warn_penalty_types(
                    update, context, query, chat_id, lang)
                return

            if action == "warn_penalty_duration":
                await CallbackHandlers._show_penalty_durations(
                    update, context, query, chat_id, lang, 'warn_penalty')
                return

            if action == "warn_toggle":
                settings = await CallbackHandlers._get_security_settings_cached(chat_id)
                new_val = 1 - _coerce_int(settings.get('warn_enabled', 0))
                await DB.update_security_settings(chat_id, warn_enabled=new_val)
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._refresh_security_view(
                    query, context, chat_id, lang)
                return

            if action == "warn_count":
                await CallbackHandlers._show_warn_count_buttons(
                    update, context, query, chat_id, lang)
                return

            if action == "penalty":
                await CallbackHandlers._show_penalty_types(
                    update, context, query, chat_id, lang)
                return

            if action == "del_pen":
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('ban_btn', lang, "🚫"),
                        callback_data=f"sec_set_del_penalty:ban:{chat_id}"),
                     InlineKeyboardButton(
                        await _trans('mute_btn', lang, "🔇"),
                        callback_data=f"sec_set_del_penalty:mute:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('kick_btn', lang, "👢"),
                        callback_data=f"sec_set_del_penalty:kick:{chat_id}"),
                     InlineKeyboardButton(
                        await _trans('restrict_btn', lang, "🔒"),
                        callback_data=f"sec_set_del_penalty:restrict:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('no_penalty_btn', lang, "🚫"),
                        callback_data=f"sec_set_del_penalty:none:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('penalty_duration_btn', lang, "⏱️"),
                        callback_data=f"sec_set_del_penalty_duration:{chat_id}")],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=f"{CB.GRP_SET}:{chat_id}")],
                ])
                await safe_edit(query,
                    await _trans('choose_delete_penalty', lang, "🚫"),
                    reply_markup=kb, bot=context.bot)
                return

            if action == "banned_words":
                await CallbackHandlers._show_banned_words_menu(
                    update, context, query, chat_id, lang)
                return

            if action == "toggle_banned_words":
                settings = await CallbackHandlers._get_security_settings_cached(chat_id)
                new_val = 1 - _coerce_int(settings.get('delete_banned_words', 0))
                await DB.update_security_settings(chat_id, delete_banned_words=new_val)
                await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                await CallbackHandlers._show_banned_words_menu(
                    update, context, query, chat_id, lang)
                return

            if action in ("close", "back"):
                StateManager.clear(user_id)
                _clear_context_keys(context)
                await CallbackHandlers._show_groups_list(
                    update, context, query, user_id, lang)
                return

            if action == "antiflood_settings":
                await CallbackHandlers._show_antiflood_settings(
                    update, context, query, chat_id, lang)
                return

            if action == "night_settings":
                await CallbackHandlers._show_night_settings(
                    update, context, query, chat_id, lang)
                return

            if action == "adv_act":
                await CallbackHandlers._show_advanced_actions(
                    update, context, query, chat_id, lang)
                return

            if action == "slow_mode_seconds":
                StateManager.set(user_id, UserState.WAIT_SLOW_MODE_SECONDS)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_slow_mode_seconds', lang, "⏱️"),
                    bot=context.bot)
                return

            if action == "welcome_text":
                StateManager.set(user_id, UserState.WAIT_WELCOME_TEXT)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_welcome_text', lang, "📝"), bot=context.bot)
                return

            if action == "goodbye_text":
                StateManager.set(user_id, UserState.WAIT_GOODBYE_TEXT)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_goodbye_text', lang, "📝"), bot=context.bot)
                return

            if action == "set_antiflood_messages":
                StateManager.set(user_id, UserState.WAIT_ANTIFLOOD_MESSAGES)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_antiflood_messages', lang, "📊"),
                    bot=context.bot)
                return

            if action == "set_antiflood_seconds":
                StateManager.set(user_id, UserState.WAIT_ANTIFLOOD_SECONDS)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_antiflood_seconds', lang, "⏱️"),
                    bot=context.bot)
                return

            if action == "antiflood_penalty":
                await CallbackHandlers._show_penalty_type_selection(
                    update, context, query, chat_id, lang, 'antiflood_penalty')
                return

            if action == "set_night_start":
                StateManager.set(user_id, UserState.WAIT_NIGHT_START)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_night_start', lang, "🌙"), bot=context.bot)
                return

            if action == "set_night_end":
                StateManager.set(user_id, UserState.WAIT_NIGHT_END)
                _set_sec_chat(context, chat_id)
                await safe_edit(query,
                    await _trans('send_night_end', lang, "🌙"), bot=context.bot)
                return

            if action == "night_action":
                await CallbackHandlers._show_penalty_type_selection(
                    update, context, query, chat_id, lang, 'night_action')
                return

            if action in ("violation_settings", "violation_penalties"):
                await CallbackHandlers._show_violation_penalties(
                    update, context, query, chat_id, lang)
                return

            if action == "violation_penalty":
                await CallbackHandlers._show_penalty_type_selection(
                    update, context, query, chat_id, lang, 'violation_penalty')
                return

            logger.debug(f"⚠️ unknown sec action: {action}")
            await safe_edit(query,
                await _trans('not_available', lang, "⚠️"), bot=context.bot)

        except Exception as e:
            logger.error(f"security error: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # قناة السجل
    # =================================================================

    @staticmethod
    async def _show_log_channel_menu(query, context, chat_id, user_id, lang):
        data = await _get_log_channel_menu_data(chat_id)
        current = data.get('current')
        effective = data.get('effective')
        share_count = data.get('share_count', 0)
        share_names = data.get('share_names', [])

        if current:
            if share_count == 0:
                status_block = (f"✅ <b>Private</b>\n"
                                f"🆔 <code>{current}</code>")
            else:
                others_text = "، ".join(_html.escape(str(n)) for n in share_names)
                if share_count > 3:
                    others_text += f" +{share_count - 3}"
                status_block = (f"🤝 <b>Shared</b>\n"
                                f"🆔 <code>{current}</code>\n"
                                f"👥 {share_count + 1}\n"
                                f"📋 {others_text}")
        elif effective:
            status_block = (f"🌐 <b>Global</b>\n"
                            f"🆔 <code>{effective}</code>")
        else:
            status_block = "❌"

        help_text = KeyboardFactory.get_text("log_channel_help", lang) or ""
        title = await _trans('log_channel_btn', lang, "📢")
        text = (f"{title}\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
                f"{status_block}\n\n<i>{_html.escape(help_text)}</i>\n\n"
                f"🆔 <code>{chat_id}</code>")

        if current:
            set_text = KeyboardFactory.get_text("log_channel_change", lang) or "🔄"
        else:
            set_text = KeyboardFactory.get_text("log_channel_set", lang) or "🔗"
        test_text = KeyboardFactory.get_text("log_channel_test", lang) or "🧪"
        remove_text = KeyboardFactory.get_text("log_channel_remove", lang) or "🗑️"
        back_text = KeyboardFactory.get_text("back", lang) or "🔙"

        rows = [[InlineKeyboardButton(
            set_text, callback_data=f"log_channel_set:{chat_id}")]]
        if current:
            rows.append([
                InlineKeyboardButton(test_text,
                                     callback_data=f"log_channel_test:{chat_id}"),
                InlineKeyboardButton(remove_text,
                                     callback_data=f"log_channel_remove:{chat_id}"),
            ])
        rows.append([InlineKeyboardButton(
            back_text, callback_data=f"{CB.GRP_SET}:{chat_id}")])

        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(rows),
                        parse_mode='HTML', bot=context.bot)

    @staticmethod
    async def _handle_log_channel(update, context, query, user_id, lang):
        data = query.data or ""
        parts = data.split(":")

        chat_id = None
        if len(parts) >= 2 and parts[1].lstrip('-').isdigit():
            chat_id = int(parts[1])
        if chat_id is None:
            stored = (context.user_data.get('security_chat_id')
                      or context.user_data.get('sec_chat'))
            if stored:
                try:
                    chat_id = int(stored)
                except (TypeError, ValueError):
                    chat_id = None

        if chat_id is None:
            await safe_edit(query,
                await _trans('group_not_specified', lang, "❌"), bot=context.bot)
            return

        if not await _check_sec_auth(context, user_id, chat_id):
            await safe_edit(query,
                await _trans('no_permission', lang, "❌"), bot=context.bot)
            return

        if parts[0] == "log_channel_btn":
            action = "menu"
        elif parts[0].startswith("log_channel_"):
            action = parts[0][len("log_channel_"):]
        else:
            action = parts[0]

        try:
            if action in ("menu", "btn", "show"):
                await CallbackHandlers._show_log_channel_menu(
                    query, context, chat_id, user_id, lang)
                return

            if action == "set":
                StateManager.set(user_id, UserState.WAIT_LOG_CH)
                context.user_data['log_group_id'] = chat_id
                await safe_edit(query,
                    await _trans('log_channel_set', lang, "🔗"),
                    parse_mode='HTML', bot=context.bot)
                return

            if action == "remove":
                current = await DB.get_group_log_channel(chat_id)
                share_info = ""
                if current:
                    try:
                        others = await DB.get_groups_sharing_log_channel(
                            current, exclude_group_id=chat_id)
                        share_count = len(others) if others else 0
                        if share_count > 0:
                            share_info = f"\n\n⚠️ {share_count}"
                    except Exception:
                        pass
                ok = await DB.remove_group_log_channel(chat_id)
                await _invalidate_log_channel_menu_cache(chat_id)
                if not ok:
                    await safe_edit(query,
                        await _trans('delete_failed', lang, "❌"), bot=context.bot)
                    return
                back_text = KeyboardFactory.get_text("back", lang) or "🔙"
                msg = await _trans('log_channel_removed', lang, "🗑️")
                await safe_edit(
                    query,
                    f"{msg}\n\n🆔 <code>{chat_id}</code>{share_info}",
                    reply_markup=InlineKeyboardMarkup([[
                        InlineKeyboardButton(back_text,
                                             callback_data=f"{CB.GRP_SET}:{chat_id}")]]),
                    parse_mode='HTML', bot=context.bot)
                return

            if action == "test":
                current = await DB.get_group_log_channel(chat_id)
                if not current:
                    current = await DB.get_log_channel()
                if not current:
                    await safe_edit(query,
                        await _trans('log_channel_none', lang, "❌"),
                        bot=context.bot)
                    return
                try:
                    test_msg = await _trans('log_channel_test', lang, "🧪")
                    await safe_send(
                        context.bot, current,
                        f"{test_msg}\n"
                        f"🆔 <code>{chat_id}</code>\n"
                        f"🕐 {TimeUtils.mecca_iso()}",
                        parse_mode='HTML')
                    await _invalidate_log_channel_menu_cache(chat_id)
                    await safe_edit(query,
                        f"✅ <code>{current}</code>",
                        parse_mode='HTML', bot=context.bot)
                except Exception as e:
                    logger.error(f"test log: {e}", exc_info=True)
                    await safe_edit(query, f"❌ {str(e)[:100]}", bot=context.bot)
                return

            await safe_edit(query,
                await _trans('unknown_action', lang, "⚠️"), bot=context.bot)

        except Exception as e:
            logger.error(f"_handle_log_channel: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # دوال عرض الأمان
    # ═════════════════════════════════════════════════════════════

    @staticmethod
    async def _show_warn_count_buttons(update, context, query, chat_id, lang):
        settings = await CallbackHandlers._get_security_settings_cached(chat_id)
        current = _coerce_int(settings.get('max_warnings'), 3)
        title = await _trans('warn_count_title', lang, "🔢")
        current_label = _fmt(await _trans('warn_count_current', lang, "{count}"),
                             count=current)
        choose = await _trans('warn_count_choose', lang, "")
        text = f"{title}\n\n{current_label}\n\n{choose}"
        counts = [1, 2, 3, 4, 5, 10]
        kb = []
        row = []
        for n in counts:
            icon = "✅" if n == current else ""
            row.append(InlineKeyboardButton(
                f"{icon} {n}", callback_data=f"set_warn_count:{chat_id}:{n}"))
            if len(row) == 3:
                kb.append(row)
                row = []
        if row:
            kb.append(row)
        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang),
            callback_data=f"sec_warn:{chat_id}")])
        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(kb),
                        parse_mode='HTML', bot=context.bot)

    @staticmethod
    async def _show_warn_penalty_types(update, context, query, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('ban_btn', lang, "🚫"),
                callback_data=f"set_warn_penalty:ban:{chat_id}"),
             InlineKeyboardButton(
                await _trans('mute_btn', lang, "🔇"),
                callback_data=f"set_warn_penalty:mute:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('kick_btn', lang, "👢"),
                callback_data=f"set_warn_penalty:kick:{chat_id}"),
             InlineKeyboardButton(
                await _trans('restrict_btn', lang, "🔒"),
                callback_data=f"set_warn_penalty:restrict:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('choose_warn_penalty', lang, "⚖️"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_banned_words_menu(update, context, query, chat_id, lang):
        settings = await CallbackHandlers._get_security_settings_cached(chat_id)
        is_enabled = _coerce_int(settings.get('delete_banned_words'), 0)
        if is_enabled:
            toggle_text = await _trans('disable_delete', lang, "❌")
        else:
            toggle_text = await _trans('enable_delete', lang, "✅")
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('add_word', lang, "➕"),
                callback_data=f"ban_add:{chat_id}"),
             InlineKeyboardButton(
                await _trans('words_list', lang, "📋"),
                callback_data=f"ban_list:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('delete_word', lang, "🗑️"),
                callback_data=f"ban_rem:{chat_id}")],
            [InlineKeyboardButton(
                toggle_text,
                callback_data=f"sec_toggle_banned_words:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('manage_banned_words', lang, "🚫"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_penalty_type_selection(update, context, query, chat_id,
                                            lang, setting_key):
        penalty_types = [
            (await _trans('mute_btn', lang, "🔇"), "mute"),
            (await _trans('ban_btn', lang, "🚫"), "ban"),
            (await _trans('kick_btn', lang, "👢"), "kick"),
            (await _trans('restrict_btn', lang, "🔒"), "restrict"),
            (await _trans('no_penalty_btn', lang, "🚫"), "none"),
        ]
        kb = []
        for label, ptype in penalty_types:
            callback = f"sec_set_{setting_key}:{chat_id}:{ptype}"
            kb.append([InlineKeyboardButton(label, callback_data=callback)])
        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang),
            callback_data=f"{CB.GRP_SET}:{chat_id}")])
        text = await _trans('choose_penalty_type', lang, "🚫")
        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(kb),
                        bot=context.bot)

    @staticmethod
    async def _show_all_penalty_durations_menu(query, context, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('mute_duration_btn', lang, "⏱️"),
                callback_data=f"sec_set_mute_duration:{chat_id}"),
             InlineKeyboardButton(
                await _trans('ban_duration_btn', lang, "⏱️"),
                callback_data=f"sec_set_ban_duration:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('restrict_duration_btn', lang, "⏱️"),
                callback_data=f"sec_set_restrict_duration:{chat_id}"),
             InlineKeyboardButton(
                await _trans('warn_duration_btn', lang, "⏱️"),
                callback_data=f"sec_warn_penalty_duration:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('flood_duration_btn', lang, "⏱️"),
                callback_data=f"sec_antiflood_duration:{chat_id}"),
             InlineKeyboardButton(
                await _trans('night_duration_btn', lang, "⏱️"),
                callback_data=f"sec_night_duration:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('delete_penalty_duration_btn', lang, "⏱️"),
                callback_data=f"sec_set_del_penalty_duration:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('penalty_durations_title', lang, "⏱️"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_penalty_durations(update, context, query, chat_id,
                                       lang, penalty_type='mute'):
        if penalty_type == 'kick':
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton(
                    KeyboardFactory.get_text("back", lang),
                    callback_data=f"{CB.GRP_SET}:{chat_id}")]])
            await safe_edit(query,
                await _trans('kick_no_duration', lang, "✅"),
                reply_markup=kb, bot=context.bot)
            return
        durations = [
            (await _trans('permanent', lang, "∞"), 0),
            (await _trans('half_hour', lang, "30m"), 1800),
            (await _trans('hour', lang, "1h"), 3600),
            (await _trans('day', lang, "1d"), 86400),
            (await _trans('week', lang, "1w"), 604800),
            (await _trans('ten_days', lang, "10d"), 864000),
            (await _trans('month', lang, "1mo"), 2592000),
        ]
        kb = []
        for i in range(0, len(durations), 2):
            row = []
            name, secs = durations[i]
            row.append(InlineKeyboardButton(
                name, callback_data=f"set_duration:{penalty_type}:{chat_id}:{secs}"))
            if i + 1 < len(durations):
                name2, secs2 = durations[i + 1]
                row.append(InlineKeyboardButton(
                    name2,
                    callback_data=f"set_duration:{penalty_type}:{chat_id}:{secs2}"))
            kb.append(row)
        kb.append([InlineKeyboardButton(
            KeyboardFactory.get_text("back", lang),
            callback_data=f"{CB.GRP_SET}:{chat_id}")])
        type_key = f"duration_type_{penalty_type}"
        type_name = await _trans(type_key, lang, penalty_type)
        msg = _fmt(await _trans('choose_duration_for', lang,
                                "⏱️ {type_name}"), type_name=type_name)
        await safe_edit(query, msg, reply_markup=InlineKeyboardMarkup(kb),
                        bot=context.bot)

    @staticmethod
    async def _show_violation_penalties(update, context, query, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('violation_strikes_btn', lang, "🔢"),
                callback_data=f"sec_set_violation_strikes:{chat_id}"),
             InlineKeyboardButton(
                await _trans('violation_duration_btn', lang, "⏱️"),
                callback_data=f"sec_set_violation_duration:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('violation_penalty_btn', lang, "⚖️"),
                callback_data=f"sec_violation_penalty:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('violation_settings_title', lang, "🚨"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_antiflood_settings(update, context, query, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('messages_count_btn', lang, "🔢"),
                callback_data=f"sec_set_antiflood_messages:{chat_id}"),
             InlineKeyboardButton(
                await _trans('seconds_btn', lang, "⏱️"),
                callback_data=f"sec_set_antiflood_seconds:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('penalty_type_btn', lang, "🚫"),
                callback_data=f"sec_antiflood_penalty:{chat_id}"),
             InlineKeyboardButton(
                await _trans('penalty_duration_btn', lang, "⏱️"),
                callback_data=f"sec_antiflood_duration:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('antiflood_settings_title', lang, "🌊"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_night_settings(update, context, query, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('start_time_btn', lang, "🌙"),
                callback_data=f"sec_set_night_start:{chat_id}"),
             InlineKeyboardButton(
                await _trans('end_time_btn', lang, "🌙"),
                callback_data=f"sec_set_night_end:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('action_type_btn', lang, "🎬"),
                callback_data=f"sec_night_action:{chat_id}"),
             InlineKeyboardButton(
                await _trans('action_duration_btn', lang, "⏱️"),
                callback_data=f"sec_night_duration:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('night_mode_settings', lang, "🌙"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_advanced_actions(update, context, query, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('deactivate_all_btn', lang, "🔴"),
                callback_data=f"sec_deactivate_all:{chat_id}"),
             InlineKeyboardButton(
                await _trans('activate_all_btn_short', lang, "🟢"),
                callback_data=f"sec_activate_all:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('act_ban', lang, "🚫"),
                callback_data=f"act_ban:{chat_id}"),
             InlineKeyboardButton(
                await _trans('act_mute', lang, "🔇"),
                callback_data=f"act_mute:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('act_kick', lang, "👢"),
                callback_data=f"act_kick:{chat_id}"),
             InlineKeyboardButton(
                await _trans('act_restrict', lang, "🔒"),
                callback_data=f"act_restrict:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('act_unban', lang, "🔓"),
                callback_data=f"act_unban:{chat_id}"),
             InlineKeyboardButton(
                await _trans('act_warn', lang, "⚠️"),
                callback_data=f"act_warn:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('act_pin', lang, "📌"),
                callback_data=f"act_pin:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('log_btn', lang, "📋"),
                callback_data=f"act_log:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('advanced_actions_title', lang, "🛠️"),
            reply_markup=kb, bot=context.bot)

    @staticmethod
    async def _show_admin_logs(update, context, query, chat_id, lang):
        logs = await DB.get_admin_logs(chat_id, 10)
        if logs:
            lines = []
            for l in logs:
                ld = _row_to_dict(l) or {}
                lines.append(f"• {_safe_str(ld.get('admin_id'))} → "
                             f"{_safe_str(ld.get('action'))}")
            text = await _trans('admin_logs_title', lang, "📋") + "\n\n" + "\n".join(lines)
        else:
            text = await _trans('no_data', lang, "📭")
        await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
            [[InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")]]),
            bot=context.bot)

    @staticmethod
    async def _show_penalty_types(update, context, query, chat_id, lang):
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(
                await _trans('ban_btn', lang, "🚫"),
                callback_data=f"sec_penalty_ban:{chat_id}"),
             InlineKeyboardButton(
                await _trans('mute_btn', lang, "🔇"),
                callback_data=f"sec_penalty_mute:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('kick_btn', lang, "👢"),
                callback_data=f"sec_penalty_kick:{chat_id}"),
             InlineKeyboardButton(
                await _trans('restrict_btn', lang, "🔒"),
                callback_data=f"sec_penalty_restrict:{chat_id}")],
            [InlineKeyboardButton(
                await _trans('no_penalty_btn', lang, "🚫"),
                callback_data=f"sec_penalty_none:{chat_id}")],
            [InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang),
                callback_data=f"{CB.GRP_SET}:{chat_id}")],
        ])
        await safe_edit(query,
            await _trans('choose_penalty_type', lang, "🚫"),
            reply_markup=kb, bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Admin
    # =================================================================

    @staticmethod
    async def _handle_admin(update, context, query, user_id, lang=None):
        if not CONFIG.is_developer(user_id):
            await safe_edit(query,
                await _trans('unauthorized', lang or 'ar', "❌"),
                bot=context.bot)
            return
        if not lang:
            lang = await DB.get_user_language(user_id) or 'ar'
        data = query.data

        try:
            if data == "admin_grant_free":
                StateManager.set(user_id, UserState.WAIT_GRANT_FREE)
                await safe_edit(query,
                    await _trans('grant_usage', lang, "🎁"), bot=context.bot)
                return

            if data == "admin_ban_user":
                StateManager.set(user_id, UserState.WAIT_BAN_USER_ID)
                await safe_edit(query,
                    await _trans('ban_user_title', lang, "🚫") + "\n\n" +
                    await _trans('ban_user_prompt', lang, "") + "\n" +
                    await _trans('ban_user_example', lang, ""),
                    bot=context.bot)
                return

            if data == "admin_unban_user":
                StateManager.set(user_id, UserState.WAIT_UNBAN_USER_ID)
                await safe_edit(query,
                    await _trans('unban_user_title', lang, "✅") + "\n\n" +
                    await _trans('unban_user_prompt', lang, "") + "\n" +
                    await _trans('unban_user_example', lang, ""),
                    bot=context.bot)
                return

            if data == CB.ADMIN_USERS:
                try:
                    stats = await DB.get_user_stats() or {}
                    if not isinstance(stats, dict):
                        stats = _row_to_dict(stats) or {}
                except Exception:
                    stats = {}
                text = (await _trans('users_title', lang, "👥") + "\n\n"
                        + _fmt(await _trans('users_total', lang, "👥 {count}"),
                               count=stats.get('users', 0)) + "\n"
                        + _fmt(await _trans('users_banned', lang, "⛔ {count}"),
                               count=stats.get('banned', 0)))
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('banned_users_btn', lang, "⛔"),
                        callback_data=CB.ADMIN_BANNED)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_BANNED:
                banned_users = await DB.fetchall(
                    "SELECT user_id FROM users WHERE banned=1 LIMIT 20")
                if banned_users:
                    text = (await _trans('banned_users_title', lang, "⛔") + "\n\n"
                            + "\n".join(
                                _safe_str((_row_to_dict(u) or {}).get('user_id'))
                                for u in banned_users))
                else:
                    text = await _trans('no_banned_users', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('unban_all_btn', lang, "✅"),
                        callback_data=CB.ADMIN_UNBAN_ALL)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_UNBAN_ALL:
                await DB.execute("UPDATE users SET banned=0 WHERE banned=1")
                await safe_edit(query,
                    await _trans('unban_all_success', lang, "✅"), bot=context.bot)
                return

            if data == CB.ADMIN_STATS:
                try:
                    stats = await DB.get_general_stats() or {}
                    if not isinstance(stats, dict):
                        stats = _row_to_dict(stats) or {}
                except Exception:
                    stats = {}
                text = (await _trans('general_stats_title', lang, "📊") + "\n\n"
                        + _fmt(await _trans('stats_users', lang, "👥 {count}"),
                               count=stats.get('users', 0)) + "\n"
                        + _fmt(await _trans('stats_channels', lang, "📡 {count}"),
                               count=stats.get('channels', 0)) + "\n"
                        + _fmt(await _trans('stats_groups', lang, "👥 {count}"),
                               count=stats.get('groups', 0)) + "\n"
                        + _fmt(await _trans('stats_posts', lang, "📝 {count}"),
                               count=stats.get('posts', 0)) + "\n"
                        + _fmt(await _trans('stats_published', lang, "✅ {count}"),
                               count=stats.get('published', 0)) + "\n"
                        + _fmt(await _trans('stats_invoices', lang, "🧾 {count}"),
                               count=stats.get('invoices', 0)) + "\n"
                        + _fmt(await _trans('stats_tickets', lang, "🎫 {count}"),
                               count=stats.get('tickets', 0)))
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)]])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_CHANNELS:
                context.user_data['adm_ch_page'] = 0
                await CallbackHandlers._show_admin_channels(
                    update, context, query, user_id, lang)
                return

            if data.startswith("admin_toggle_ch:"):
                ch_db_id = _coerce_int(data.split(":")[-1])
                if ch_db_id <= 0:
                    await safe_edit(query,
                        await _trans('invalid_data', lang, "❌"), bot=context.bot)
                    return
                await DB.execute(
                    "UPDATE user_channels SET banned = 1 - banned WHERE id=?",
                    (ch_db_id,))
                await CallbackHandlers._show_admin_channels(
                    update, context, query, user_id, lang)
                return

            if data == CB.ADMIN_GROUPS:
                context.user_data['adm_gr_page'] = 0
                await CallbackHandlers._show_admin_groups(
                    update, context, query, user_id, lang)
                return

            if data.startswith("admin_toggle_gr:"):
                chat_id = _coerce_int(data.split(":")[-1])
                if chat_id == 0:
                    await safe_edit(query,
                        await _trans('invalid_data', lang, "❌"), bot=context.bot)
                    return
                row = await DB.fetchone(
                    "SELECT banned FROM bot_groups WHERE chat_id=?", (chat_id,))
                row_d = _row_to_dict(row) or {}
                if row_d:
                    new_val = 0 if row_d.get('banned') else 1
                    if new_val == 1:
                        try:
                            await context.bot.leave_chat(chat_id)
                        except Exception:
                            pass
                    await DB.execute(
                        "UPDATE bot_groups SET banned=? WHERE chat_id=?",
                        (new_val, chat_id))
                await CallbackHandlers._show_admin_groups(
                    update, context, query, user_id, lang)
                return

            if data == CB.ADMIN_ADD_ADMIN:
                StateManager.set(user_id, UserState.WAIT_ADMIN_ADD)
                await safe_edit(query,
                    await _trans('add_admin_prompt', lang, "👑"), bot=context.bot)
                return

            if data == CB.ADMIN_REM_ADMIN:
                StateManager.set(user_id, UserState.WAIT_ADMIN_REM)
                await safe_edit(query,
                    await _trans('remove_admin_prompt', lang, "🗑️"), bot=context.bot)
                return

            if data == CB.ADMIN_LIST_ADMINS:
                admins = await DB.get_admin_list()
                if admins:
                    text = (await _trans('admins_list_title', lang, "👑") + "\n\n"
                            + "\n".join(
                                f"• {_safe_str((_row_to_dict(a) or {}).get('user_id'))}"
                                for a in admins))
                else:
                    text = await _trans('no_admins', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('add_btn', lang, "➕"),
                        callback_data=CB.ADMIN_ADD_ADMIN),
                     InlineKeyboardButton(
                        await _trans('remove_btn', lang, "🗑️"),
                        callback_data=CB.ADMIN_REM_ADMIN)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_BROADCAST:
                StateManager.set(user_id, UserState.WAIT_BROADCAST)
                await safe_edit(query,
                    await _trans('broadcast_prompt', lang, "📨"), bot=context.bot)
                return

            if data == CB.ADMIN_INVOICES:
                invoices = await DB.fetchall(
                    "SELECT number, amount, status FROM invoices "
                    "ORDER BY id DESC LIMIT 20")
                if invoices:
                    lines = []
                    for i in invoices:
                        id_ = _row_to_dict(i) or {}
                        lines.append(
                            f"• {_safe_str(id_.get('number'))} - "
                            f"{_safe_str(id_.get('amount'), '0')} ⭐ - "
                            f"{_safe_str(id_.get('status'))}")
                    text = await _trans('invoices_title', lang, "🧾") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_invoices_admin', lang, "📭")
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)]])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_BACKUP:
                await safe_edit(query,
                    await _trans('backup_started', lang, "⏳"), bot=context.bot)
                task = asyncio.create_task(
                    CallbackHandlers._do_backup(context, user_id, lang))
                ACTIVE_TASKS.add(task)
                task.add_done_callback(ACTIVE_TASKS.discard)
                return

            if data in (CB.ADMIN_RESTORE, CB.ADMIN_RESTORE_SEL):
                await CallbackHandlers._show_restore_backups(
                    update, context, query, user_id, lang)
                return

            if data.startswith("admin_restore_file:"):
                fname = data.split(":", 1)[1]
                backup_file = PATHS.BACKUPS / fname
                try:
                    resolved = backup_file.resolve(strict=False)
                    base = PATHS.BACKUPS.resolve(strict=False)
                except OSError:
                    await safe_edit(query, "❌", bot=context.bot)
                    return
                if resolved.parent != base or not backup_file.exists():
                    await safe_edit(query, "❌", bot=context.bot)
                    return
                try:
                    pre_restore = (PATHS.BACKUPS / f"pre_restore_"
                                   f"{TimeUtils.mecca_now().strftime('%Y%m%d_%H%M%S_%f')}.db")
                    shutil.copy2(PATHS.DB, pre_restore)
                    db_closed = False
                    try:
                        close_fn = getattr(DB, 'close', None)
                        if callable(close_fn):
                            await close_fn()
                            db_closed = True
                    except Exception:
                        pass
                    shutil.copy2(backup_file, PATHS.DB)
                    if db_closed:
                        try:
                            reconnect_fn = getattr(DB, 'reconnect', None)
                            if callable(reconnect_fn):
                                await reconnect_fn()
                            else:
                                init_fn = getattr(DB, 'initialize_db', None)
                                if callable(init_fn):
                                    await init_fn()
                                else:
                                    init_fn = getattr(DB, 'initialize', None)
                                    if callable(init_fn):
                                        await init_fn()
                        except Exception:
                            pass
                    await safe_edit(query,
                        await _trans('restore_success', lang, "✅"),
                        bot=context.bot)
                except Exception as e:
                    await safe_edit(query, f"❌ {str(e)[:100]}", bot=context.bot)
                return

            if data == CB.ADMIN_RAM:
                ram = get_ram_usage() or {}
                text = (await _trans('ram_title', lang, "🖥️") + "\n\n"
                        + _fmt(await _trans('ram_total', lang, "💾 {value}"),
                               value=ram.get('total', 0)) + "\n"
                        + _fmt(await _trans('ram_used', lang, "📊 {value}"),
                               value=ram.get('used', 0)) + "\n"
                        + _fmt(await _trans('ram_percent', lang, "📈 {value}"),
                               value=ram.get('percent', 0)))
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_METRICS:
                try:
                    stats = await DB.get_general_stats() or {}
                    if not isinstance(stats, dict):
                        stats = _row_to_dict(stats) or {}
                except Exception:
                    stats = {}
                try:
                    db_size_kb = float(stats.get('db_size_kb', 0) or 0)
                    if db_size_kb == 0:
                        db_size_kb = await DB.get_db_size_kb()
                except Exception:
                    db_size_kb = 0.0
                if db_size_kb >= 1024 * 1024:
                    size_display = f"{db_size_kb / (1024 * 1024):.2f} GB"
                elif db_size_kb >= 1024:
                    size_display = f"{db_size_kb / 1024:.2f} MB"
                else:
                    size_display = f"{db_size_kb:.1f} KB"
                text = (await _trans('metrics_title', lang, "📊") + "\n\n"
                        + _fmt(await _trans('stats_users', lang, "👥 {count}"),
                               count=stats.get('users', 0)) + "\n"
                        + _fmt(await _trans('stats_channels', lang, "📡 {count}"),
                               count=stats.get('channels', 0)) + "\n"
                        + _fmt(await _trans('stats_groups', lang, "👥 {count}"),
                               count=stats.get('groups', 0)) + "\n"
                        + _fmt(await _trans('stats_posts', lang, "📝 {count}"),
                               count=stats.get('posts', 0)) + "\n"
                        + _fmt(await _trans('stats_published', lang, "✅ {count}"),
                               count=stats.get('published', 0)) + "\n"
                        + _fmt(await _trans('stats_invoices', lang, "🧾 {count}"),
                               count=stats.get('invoices', 0)) + "\n"
                        + _fmt(await _trans('stats_tickets', lang, "🎫 {count}"),
                               count=stats.get('tickets', 0)) + "\n"
                        + _fmt(await _trans('metrics_db_size', lang, "💾 {size}"),
                               size=size_display))
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_UPTIME:
                uptime = (time.monotonic()
                          - context.bot_data.get('start_time', time.monotonic()))
                hours, remainder = divmod(uptime, 3600)
                minutes, seconds = divmod(remainder, 60)
                text = _fmt(await _trans('uptime_format', lang, "{h}h {m}m {s}s"),
                            h=int(hours), m=int(minutes), s=int(seconds))
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_TICKETS:
                tickets = await DB.get_tickets()
                if tickets:
                    lines = []
                    for t in tickets[:10]:
                        td = _row_to_dict(t) or {}
                        lines.append(
                            f"• #{_safe_str(td.get('ticket_number'))} - "
                            f"{_safe_str(td.get('user_id'))}: "
                            f"{_safe_str(td.get('message'), '')[:50]}")
                    text = await _trans('tickets_title', lang, "🎫") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_tickets', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('delete_all_btn', lang, "🗑️"),
                        callback_data=CB.ADMIN_DEL_TICKETS)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_DEL_TICKETS:
                await DB.delete_all_tickets()
                await safe_edit(query,
                    await _trans('tickets_deleted', lang, "✅"), bot=context.bot)
                return

            if data == CB.ADMIN_PAYMENT_LOGS:
                logs = await DB.fetchall(
                    "SELECT user_id, event_type, created_at "
                    "FROM payment_logs ORDER BY id DESC LIMIT 20")
                if logs:
                    lines = []
                    for l in logs:
                        ld = _row_to_dict(l) or {}
                        lines.append(
                            f"• {_safe_str(ld.get('user_id'))} - "
                            f"{_safe_str(ld.get('event_type'))} "
                            f"({_safe_str(ld.get('created_at'))})")
                    text = await _trans('payment_logs_title', lang, "💳") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_data', lang, "📭")
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_SET_UPDATE_CH:
                StateManager.set(user_id, UserState.WAIT_UPDATE_CH)
                await safe_edit(query,
                    await _trans('send_update_ch_prompt', lang, "📢"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_SEND_UPDATE:
                StateManager.set(user_id, UserState.WAIT_UPDATE)
                await safe_edit(query,
                    await _trans('send_update_prompt', lang, "📝"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_SHOW_UPDATE:
                ch = await DB.get_updates_channel()
                if ch:
                    text = _fmt(await _trans('update_channel_set', lang, "📢 {ch}"),
                                ch=ch)
                else:
                    text = await _trans('update_channel_not_set', lang, "📭")
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_SET_LOG_CH:
                StateManager.set(user_id, UserState.WAIT_LOG_CH)
                await safe_edit(query,
                    await _trans('send_log_ch_prompt', lang, "📋"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_LOG_CH:
                ch = await DB.get_log_channel()
                if ch:
                    text = _fmt(await _trans('log_channel_set', lang, "📋 {ch}"),
                                ch=ch)
                else:
                    text = await _trans('log_channel_not_set', lang, "📭")
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_FORCE_SUB:
                sub = await DB.get_force_subscribe_channel()
                status = (await _trans('force_sub_enabled', lang, "✅")
                          if sub else await _trans('force_sub_disabled', lang, "❌"))
                text = _fmt(await _trans('force_sub_status', lang, "🔒 {status}"),
                            status=status)
                if sub:
                    text += "\n" + _fmt(
                        await _trans('force_sub_channel', lang, "Channel: {ch}"),
                        ch=sub)
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_SET_FORCE:
                StateManager.set(user_id, UserState.WAIT_FORCE)
                await safe_edit(query,
                    await _trans('send_force_ch_prompt', lang, "🔒"),
                    bot=context.bot)
                return

            if data == "admin_disable_force":
                await DB.execute(
                    "UPDATE settings SET value = '' "
                    "WHERE key = 'force_subscribe_channel'")
                try:
                    _invalidate_force_sub_cache()
                except Exception:
                    pass
                await safe_edit(query,
                    await _trans('force_disabled_success', lang, "✅"),
                    bot=context.bot)
                return

            if data == "admin_upload_backup":
                StateManager.set(user_id, UserState.WAIT_BACKUP_FILE)
                await safe_edit(query,
                    await _trans('upload_backup_prompt', lang, "📤"),
                    bot=context.bot)
                return

            if data == "admin_show_backups":
                await CallbackHandlers._show_restore_backups(
                    update, context, query, user_id, lang)
                return

            if data == CB.ADMIN_REFRESH_CACHE:
                _invalidate_sec_auth_cache()
                try:
                    await invalidate_user_cache(user_id)
                except Exception:
                    pass
                try:
                    await _security_settings_cache.clear()
                    await _security_stats_cache_local.clear()
                except Exception:
                    pass
                await safe_edit(query,
                    await _trans('cache_refreshed_admin', lang, "🔄"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_BANNED_CH:
                banned_channels = await DB.fetchall(
                    "SELECT channel_id, channel_name FROM user_channels "
                    "WHERE banned=1 LIMIT 20")
                if banned_channels:
                    lines = []
                    for c in banned_channels:
                        cd = _row_to_dict(c) or {}
                        lines.append(
                            f"• {_safe_str(cd.get('channel_name'))} "
                            f"({_safe_str(cd.get('channel_id'))})")
                    text = await _trans('banned_channels_title', lang, "🚫") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_data', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('activate_all_btn', lang, "✅"),
                        callback_data=CB.ADMIN_ACTIVATE_CH)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_ACTIVATE_CH:
                await DB.execute(
                    "UPDATE user_channels SET banned=0 WHERE banned=1")
                await safe_edit(query,
                    await _trans('channels_activated', lang, "✅"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_BANNED_GR:
                banned_groups = await DB.fetchall(
                    "SELECT chat_id, chat_name FROM bot_groups "
                    "WHERE banned=1 LIMIT 20")
                if banned_groups:
                    lines = []
                    for g in banned_groups:
                        gd = _row_to_dict(g) or {}
                        lines.append(
                            f"• {_safe_str(gd.get('chat_name'))} "
                            f"({_safe_str(gd.get('chat_id'))})")
                    text = await _trans('banned_groups_title', lang, "🚫") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_data', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('unban_all_groups_btn', lang, "🔓"),
                        callback_data=CB.ADMIN_UNBAN_GR)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == CB.ADMIN_UNBAN_GR:
                await DB.execute(
                    "UPDATE bot_groups SET banned=0 WHERE banned=1")
                await safe_edit(query,
                    await _trans('groups_unbanned', lang, "✅"), bot=context.bot)
                return

            if data == CB.ADMIN_REPLIES:
                replies = await DB.fetchall(
                    "SELECT keyword FROM auto_replies "
                    "WHERE chat_id=-1 LIMIT 30")
                if replies:
                    lines = []
                    for r in replies:
                        rd = _row_to_dict(r) or {}
                        lines.append(f"• {_safe_str(rd.get('keyword'))}")
                    text = await _trans('public_replies_title', lang, "💬") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_replies', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('add_btn', lang, "➕"),
                        callback_data="admin_add_reply"),
                     InlineKeyboardButton(
                        await _trans('remove_btn', lang, "🗑️"),
                        callback_data="admin_del_reply")],
                    [InlineKeyboardButton(
                        await _trans('export_btn', lang, "📤"),
                        callback_data=CB.ADMIN_EXPORT_REPLIES),
                     InlineKeyboardButton(
                        await _trans('import_btn', lang, "📥"),
                        callback_data=CB.ADMIN_IMPORT_REPLIES)],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == "admin_add_reply":
                StateManager.set(user_id, UserState.WAIT_KEYWORD)
                context.user_data['auto_chat'] = -1
                await safe_edit(query,
                    await _trans('send_keyword_prompt', lang, "📝"),
                    bot=context.bot)
                return

            if data == "admin_del_reply":
                StateManager.set(user_id, UserState.WAIT_AUTO_DEL)
                context.user_data['auto_chat'] = -1
                await safe_edit(query,
                    await _trans('send_keyword_delete_prompt', lang, "🗑️"),
                    bot=context.bot)
                return

            if data == "admin_list_replies":
                replies = await DB.fetchall(
                    "SELECT keyword FROM auto_replies "
                    "WHERE chat_id=-1 LIMIT 50")
                if replies:
                    lines = []
                    for r in replies:
                        rd = _row_to_dict(r) or {}
                        lines.append(f"• {_safe_str(rd.get('keyword'))}")
                    text = await _trans('public_replies_list_title', lang, "📋") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_replies', lang, "📭")
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_EXPORT_REPLIES:
                try:
                    file_path = await DB.export_auto_replies_to_file()
                except (AttributeError, Exception):
                    file_path = None
                if file_path:
                    try:
                        with open(file_path, 'rb') as f:
                            await context.bot.send_document(
                                chat_id=user_id, document=f,
                                filename=Path(file_path).name)
                    except Exception as e:
                        await safe_send(context.bot, user_id, f"❌ {e}")
                    finally:
                        try:
                            os.remove(file_path)
                        except OSError:
                            pass
                else:
                    await safe_edit(query,
                        await _trans('no_replies', lang, "📭"), bot=context.bot)
                return

            if data == CB.ADMIN_IMPORT_REPLIES:
                StateManager.set(user_id, UserState.WAIT_IMPORT_FILE)
                await safe_edit(query,
                    await _trans('send_json_prompt', lang, "📤"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_IMPORT_GITHUB:
                StateManager.set(user_id, UserState.WAIT_GITHUB_URL)
                await safe_edit(query,
                    await _trans('send_url_prompt', lang, "📥"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_BANNED_WORDS:
                words = await DB.get_banned_words(-1)
                if words:
                    text = (await _trans('public_banned_words_title_full', lang, "🚫")
                            + "\n\n"
                            + "\n".join(f"• {w}" for w in words[:30]))
                else:
                    text = await _trans('no_data', lang, "📭")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('add_btn', lang, "➕"),
                        callback_data="admin_add_banned"),
                     InlineKeyboardButton(
                        await _trans('remove_btn', lang, "🗑️"),
                        callback_data="admin_rem_banned")],
                    [InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=CB.ADMIN)],
                ])
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if data == "admin_add_banned":
                StateManager.set(user_id, UserState.WAIT_GLOBAL_BAN)
                context.user_data['ban_chat'] = -1
                await safe_edit(query,
                    await _trans('send_keyword_prompt', lang, "📝"),
                    bot=context.bot)
                return

            if data == "admin_rem_banned":
                StateManager.set(user_id, UserState.WAIT_REM_GLOBAL_BAN)
                context.user_data['ban_chat'] = -1
                await safe_edit(query,
                    await _trans('send_keyword_delete_prompt', lang, "🗑️"),
                    bot=context.bot)
                return

            if data == "admin_list_banned":
                words = await DB.get_banned_words(-1)
                if words:
                    text = (await _trans('public_banned_words_list_full', lang, "📋")
                            + "\n\n"
                            + "\n".join(f"• {w}" for w in words))
                else:
                    text = await _trans('no_data', lang, "📭")
                await safe_edit(query, text, bot=context.bot)
                return

            if data == CB.ADMIN_CREATE_CONTEST:
                StateManager.set(user_id, UserState.WAIT_CONTEST_TITLE)
                await safe_edit(query,
                    await _trans('contest_title_prompt', lang, "🏆"),
                    bot=context.bot)
                return

            if data == CB.ADMIN_DECLARE_WINNER:
                contests = await DB.get_active_contests(5)
                if not contests:
                    await safe_edit(query,
                        await _trans('no_active_contests_admin', lang, "📭"),
                        bot=context.bot)
                    return
                kb = []
                for c in contests:
                    cd = _row_to_dict(c) or {}
                    kb.append([InlineKeyboardButton(
                        f"🏆 {_safe_str(cd.get('title'))[:20]}",
                        callback_data=(
                            f"{CB.DECLARE_WINNER_SEL}:{cd.get('id', 0)}"))])
                kb.append([InlineKeyboardButton(
                    KeyboardFactory.get_text("back", lang),
                    callback_data=CB.ADMIN)])
                await safe_edit(query,
                    await _trans('choose_contest', lang, "🏆"),
                    reply_markup=InlineKeyboardMarkup(kb), bot=context.bot)
                return

            if data == CB.ADMIN_DEL_CONTEST:
                contests = await DB.fetchall(
                    "SELECT id, title FROM contests "
                    "WHERE status='active' LIMIT 10")
                if not contests:
                    await safe_edit(query,
                        await _trans('no_contests_admin', lang, "📭"),
                        bot=context.bot)
                    return
                kb = []
                for c in contests:
                    cd = _row_to_dict(c) or {}
                    kb.append([InlineKeyboardButton(
                        f"🗑️ {_safe_str(cd.get('title'))[:20]}",
                        callback_data=(
                            f"admin_delete_contest:{cd.get('id', 0)}"))])
                kb.append([InlineKeyboardButton(
                    KeyboardFactory.get_text("back", lang),
                    callback_data=CB.ADMIN)])
                await safe_edit(query,
                    await _trans('choose_contest_delete', lang, "🗑️"),
                    reply_markup=InlineKeyboardMarkup(kb), bot=context.bot)
                return

            if data.startswith("admin_delete_contest:"):
                contest_id = _coerce_int(data.split(":")[-1])
                if await DB.delete_contest(contest_id, user_id):
                    await safe_edit(query,
                        await _trans('contest_deleted', lang, "✅"),
                        bot=context.bot)
                else:
                    await safe_edit(query,
                        await _trans('contest_deleted_failed', lang, "❌"),
                        bot=context.bot)
                return

            await safe_edit(query,
                await _trans('not_available', lang, "⚠️"), bot=context.bot)

        except BadRequest as e:
            if "query is too old" not in str(e).lower():
                logger.error(f"admin error: {e}", exc_info=True)
        except Exception as e:
            logger.error(f"admin error: {e}", exc_info=True)
            try:
                await safe_edit(query,
                    await _trans('error_occurred', lang, "❌"),
                    bot=context.bot)
            except Exception:
                pass

    @staticmethod
    async def _show_admin_channels(update, context, query, user_id, lang):
        try:
            total = await DB.fetchval(
                "SELECT COUNT(*) FROM user_channels", default=0)
            page = _coerce_int(context.user_data.get('adm_ch_page'), 0)
            total_pages = max(1, (total + ADMIN_PAGE_SIZE - 1) // ADMIN_PAGE_SIZE)
            if page >= total_pages:
                page = total_pages - 1
            if page < 0:
                page = 0
            context.user_data['adm_ch_page'] = page
            channels = await DB.fetchall(
                "SELECT id, channel_id, channel_name, banned "
                "FROM user_channels "
                "ORDER BY channel_name LIMIT ? OFFSET ?",
                (ADMIN_PAGE_SIZE, page * ADMIN_PAGE_SIZE))
            kb = []
            ban_l = await _trans('action_ban', lang, "🔒")
            unban_l = await _trans('action_unban', lang, "🔓")
            for c in (channels or []):
                cd = _row_to_dict(c) or {}
                action = unban_l if cd.get('banned') else ban_l
                icon = "🚫" if cd.get('banned') else "✅"
                kb.append([InlineKeyboardButton(
                    f"{icon} {_safe_str(cd.get('channel_name'))[:20]} - {action}",
                    callback_data=f"admin_toggle_ch:{cd.get('id', 0)}")])
            nav = []
            if page > 0:
                nav.append(InlineKeyboardButton(
                    await _trans('prev_btn', lang, "⬅️"),
                    callback_data=f"adm_ch_page:{page - 1}"))
            if page < total_pages - 1:
                nav.append(InlineKeyboardButton(
                    await _trans('next_btn', lang, "➡️"),
                    callback_data=f"adm_ch_page:{page + 1}"))
            if nav:
                kb.append(nav)
            kb.append([InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang), callback_data=CB.ADMIN)])
            text = (await _trans('admin_channels_title', lang, "📡")
                    + f" ({total}) — {page + 1}/{total_pages}")
            await safe_edit(query, text,
                            reply_markup=InlineKeyboardMarkup(kb),
                            bot=context.bot)
        except Exception as e:
            logger.error(f"_show_admin_channels: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _show_admin_groups(update, context, query, user_id, lang):
        try:
            total = await DB.fetchval(
                "SELECT COUNT(*) FROM bot_groups", default=0)
            page = _coerce_int(context.user_data.get('adm_gr_page'), 0)
            total_pages = max(1, (total + ADMIN_PAGE_SIZE - 1) // ADMIN_PAGE_SIZE)
            if page >= total_pages:
                page = total_pages - 1
            if page < 0:
                page = 0
            context.user_data['adm_gr_page'] = page
            groups = await DB.fetchall(
                "SELECT chat_id, chat_name, banned FROM bot_groups "
                "ORDER BY chat_name LIMIT ? OFFSET ?",
                (ADMIN_PAGE_SIZE, page * ADMIN_PAGE_SIZE))
            kb = []
            ban_l = await _trans('action_ban', lang, "🔒")
            unban_l = await _trans('action_unban', lang, "🔓")
            for g in (groups or []):
                gd = _row_to_dict(g) or {}
                action = unban_l if gd.get('banned') else ban_l
                icon = "🚫" if gd.get('banned') else "✅"
                kb.append([InlineKeyboardButton(
                    f"{icon} {_safe_str(gd.get('chat_name'))[:20]} - {action}",
                    callback_data=f"admin_toggle_gr:{gd.get('chat_id', 0)}")])
            nav = []
            if page > 0:
                nav.append(InlineKeyboardButton(
                    await _trans('prev_btn', lang, "⬅️"),
                    callback_data=f"adm_gr_page:{page - 1}"))
            if page < total_pages - 1:
                nav.append(InlineKeyboardButton(
                    await _trans('next_btn', lang, "➡️"),
                    callback_data=f"adm_gr_page:{page + 1}"))
            if nav:
                kb.append(nav)
            kb.append([InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang), callback_data=CB.ADMIN)])
            text = (await _trans('admin_groups_title', lang, "👥")
                    + f" ({total}) — {page + 1}/{total_pages}")
            await safe_edit(query, text,
                            reply_markup=InlineKeyboardMarkup(kb),
                            bot=context.bot)
        except Exception as e:
            logger.error(f"_show_admin_groups: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _show_restore_backups(update, context, query, user_id, lang='ar'):
        def _safe_mtime(p):
            try:
                return p.stat().st_mtime
            except (OSError, FileNotFoundError):
                return 0
        try:
            backups = sorted(PATHS.BACKUPS.glob("backup_*.db"),
                             key=_safe_mtime, reverse=True)
            if not backups:
                await safe_edit(query,
                    await _trans('no_backups_available', lang, "📭"),
                    bot=context.bot)
                return
            kb = []
            for b in backups[:10]:
                fname = b.name
                kb.append([InlineKeyboardButton(
                    f"📁 {fname}",
                    callback_data=f"admin_restore_file:{fname}")])
            kb.append([InlineKeyboardButton(
                KeyboardFactory.get_text("back", lang), callback_data=CB.ADMIN)])
            await safe_edit(query,
                await _trans('choose_backup_restore', lang, "📂"),
                reply_markup=InlineKeyboardMarkup(kb), bot=context.bot)
        except Exception as e:
            logger.error(f"_show_restore_backups: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Analytics
    # =================================================================

    @staticmethod
    async def _show_analytics_menu(query, context, user_id, lang):
        text = await _trans('analytics_title', lang, "📊")
        kb = None
        try:
            rows = KeyboardFactory.get_menu("analytics", lang=lang)
            if rows:
                kb = KeyboardFactory.build("analytics", lang=lang)
        except Exception as e:
            logger.warning(f"analytics build: {e}")
        if kb is None:
            kb = InlineKeyboardMarkup([[
                InlineKeyboardButton(
                    KeyboardFactory.get_text("back", lang),
                    callback_data=CB.ADMIN)]])
        await safe_edit(query, text, reply_markup=kb,
                        parse_mode='HTML', bot=context.bot)

    @staticmethod
    async def _handle_analytics(update, context, query, user_id, lang, data):
        if data.startswith("analytics_"):
            raw_action = data[len("analytics_"):]
        else:
            raw_action = data
        action = _ANALYTICS_ALIASES.get(raw_action, raw_action)

        back_btn = KeyboardFactory.get_text("back", lang)

        try:
            if action == "user_growth":
                rows = await DB.get_user_growth(30)
                if not rows:
                    await safe_edit(
                        query,
                        await _trans('growth_title', lang, "📈") + "\n\n" +
                        await _trans('no_data_enough', lang, "📭"),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        parse_mode='HTML', bot=context.bot)
                    return
                max_cnt = max(r['count'] for r in rows) or 1
                text = await _trans('growth_title', lang, "📈") + "\n"
                text += "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                total = sum(r['count'] for r in rows)
                avg = total / len(rows) if rows else 0
                text += _fmt(await _trans('growth_total', lang, "{total}"),
                             total=total) + "\n"
                text += _fmt(await _trans('growth_avg', lang, "{avg}"),
                             avg=f"{avg:.1f}") + "\n"
                text += _fmt(await _trans('growth_days_count', lang, "{count}"),
                             count=len(rows)) + "\n\n"
                text += await _trans('growth_last_10_days', lang, "") + "\n"
                for r in rows[-10:]:
                    bar_len = int((r['count'] / max_cnt) * 10)
                    bar = "█" * bar_len + "░" * (10 - bar_len)
                    date_short = r['date'][5:]
                    text += f"<code>{date_short}</code> {bar} {r['count']}\n"
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(back_btn,
                                         callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "top_channels":
                rows = await DB.get_top_channels(10)
                if not rows:
                    await safe_edit(
                        query,
                        await _trans('top_channels_title', lang, "🏆") + "\n\n" +
                        await _trans('no_channels_data', lang, "📭"),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        parse_mode='HTML', bot=context.bot)
                    return
                text = await _trans('top_channels_title', lang, "🏆") + "\n"
                text += "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                for i, ch in enumerate(rows[:10], 1):
                    rate = ch.get('success_rate', 100)
                    color = color_emoji(rate,
                        (SUCCESS_RATE_WARN_THRESHOLD, SUCCESS_RATE_GOOD_THRESHOLD))
                    name = _html.escape(str(ch['name'])[:25])
                    text += (f"{i}. {color} <b>{name}</b>\n"
                             f"   📝 {ch['total']} | ✅ {ch['published']} "
                             f"| ❌ {ch['failed']} | 🎯 {rate}%\n\n")
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(back_btn,
                                         callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "publish_stats":
                stats = await DB.get_publish_stats()
                total_ch = stats['total_channels']
                avg_posts = stats['avg_posts_per_channel']
                rate = stats.get('success_rate', DEFAULT_SUCCESS_RATE)
                completion = stats.get('completion_rate', 0)
                rate_color = color_emoji(rate,
                    (SUCCESS_RATE_WARN_THRESHOLD, SUCCESS_RATE_GOOD_THRESHOLD))
                completion_color = color_emoji(completion,
                    (SUCCESS_RATE_WARN_THRESHOLD, SUCCESS_RATE_GOOD_THRESHOLD))
                avg_color = "🟢" if avg_posts >= 20 else ("🟡" if avg_posts >= 10 else "🔴")
                attempted = stats.get('attempted',
                                      stats['published'] + stats['failed'])
                pending = stats.get('pending',
                                    max(0, stats['total_posts'] - attempted))
                text = (await _trans('publish_stats_title', lang, "📊") + "\n"
                        + "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                        + _fmt(await _trans('publish_stats_channels', lang, "{count}"),
                               count=total_ch) + "\n"
                        + _fmt(await _trans('publish_stats_total', lang, "{count}"),
                               count=stats['total_posts']) + "\n"
                        + _fmt(await _trans('publish_stats_published', lang, "{count}"),
                               count=stats['published']) + "\n"
                        + _fmt(await _trans('publish_stats_failed', lang, "{count}"),
                               count=stats['failed']) + "\n"
                        + _fmt(await _trans('publish_stats_pending', lang, "{count}"),
                               count=pending) + "\n\n"
                        + f"{avg_color} "
                        + _fmt(await _trans('publish_stats_avg_posts', lang, "{value}"),
                               value=avg_posts) + "\n"
                        + f"{avg_color} "
                        + _fmt(await _trans('publish_stats_avg_published', lang, "{value}"),
                               value=stats['avg_published_per_channel']) + "\n\n"
                        + f"{rate_color} "
                        + _fmt(await _trans('publish_stats_success_rate', lang,
                                            "{rate}% ({pub}/{att})"),
                               rate=rate, pub=stats['published'], att=attempted) + "\n"
                        + f"{completion_color} "
                        + _fmt(await _trans('publish_stats_completion_rate', lang,
                                            "{rate}% ({pub}/{total})"),
                               rate=completion, pub=stats['published'],
                               total=stats['total_posts']) + "\n")
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(back_btn,
                                         callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "channels_rate":
                rows = await DB.get_top_channels(20)
                if not rows:
                    await safe_edit(
                        query,
                        await _trans('channels_rate_title', lang, "🎯") + "\n\n"
                        + await _trans('no_data', lang, "📭"),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        parse_mode='HTML', bot=context.bot)
                    return
                text = await _trans('channels_rate_title', lang, "🎯") + "\n"
                text += "━━━━━━━━━━━━━━━━━━━━━━\n"
                text += await _trans('channels_rate_success_definition', lang, "") + "\n"
                text += await _trans('channels_rate_completion_definition', lang, "") + "\n\n"

                def _sort_key(ch):
                    sr = _coerce_float(ch.get('success_rate'), DEFAULT_SUCCESS_RATE)
                    cr = _coerce_float(ch.get('completion_rate'), 0.0)
                    return (sr, cr)

                sorted_rows = sorted(rows, key=_sort_key)
                for ch in sorted_rows[:10]:
                    text += _format_channel_rate_line(ch)
                text += await _trans('channels_rate_less_success_hint', lang, "")
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('refresh_btn', lang, "🔄"),
                        callback_data="analytics_channels_rate")],
                    [InlineKeyboardButton(back_btn,
                                          callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "subscriptions":
                rows = await DB.get_subscription_rate(6)
                if not rows:
                    await safe_edit(
                        query,
                        await _trans('subscriptions_title', lang, "💎") + "\n\n"
                        + await _trans('no_data', lang, "📭"),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        parse_mode='HTML', bot=context.bot)
                    return
                text = await _trans('subscriptions_title', lang, "💎") + "\n"
                text += "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                max_cnt = max(r['count'] for r in rows) or 1
                total = sum(r['count'] for r in rows)
                for r in rows:
                    bar_len = int((r['count'] / max_cnt) * 12)
                    bar = "█" * bar_len + "░" * (12 - bar_len)
                    text += f"<code>{r['month']}</code> {bar} {r['count']}\n"
                text += "\n"
                text += _fmt(await _trans('subscriptions_total', lang, "{count}"),
                             count=total) + "\n"
                avg = total / len(rows) if rows else 0
                text += _fmt(await _trans('subscriptions_avg', lang, "{value}"),
                             value=f"{avg:.1f}") + "\n"
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(back_btn,
                                         callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "pool":
                pool_data = await DB.get_pool_live()
                if not pool_data.get('available'):
                    await safe_edit(
                        query,
                        await _trans('pool_title', lang, "🚀") + "\n\n"
                        + _fmt(await _trans('pool_unavailable', lang, "{type}"),
                               type=pool_data.get('type', '?')),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        parse_mode='HTML', bot=context.bot)
                    return
                util = pool_data['utilization_pct']
                if util >= 80:
                    alert = await _trans('pool_alert_critical', lang, "🔴")
                elif util >= 60:
                    alert = await _trans('pool_alert_warning', lang, "🟡")
                else:
                    alert = await _trans('pool_alert_ok', lang, "🟢")
                bar_len = int(util / 100 * 15)
                bar = "█" * bar_len + "░" * (15 - bar_len)
                text = (await _trans('pool_title', lang, "🚀") + "\n"
                        + "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                        + f"{alert}\n"
                        + f"<code>{bar}</code> {util}%\n\n"
                        + _fmt(await _trans('pool_type_label', lang, "{type}"),
                               type=pool_data['type']) + "\n"
                        + _fmt(await _trans('pool_max_label', lang, "{value}"),
                               value=pool_data['max_size']) + "\n"
                        + _fmt(await _trans('pool_current_label', lang, "{value}"),
                               value=pool_data['current_size']) + "\n"
                        + _fmt(await _trans('pool_idle_label', lang, "{value}"),
                               value=pool_data['idle_size']) + "\n"
                        + _fmt(await _trans('pool_in_use_label', lang, "{value}"),
                               value=pool_data['in_use']) + "\n")
                kb = InlineKeyboardMarkup([[
                    InlineKeyboardButton(
                        await _trans('refresh_btn', lang, "🔄"),
                        callback_data="analytics_pool"),
                    InlineKeyboardButton(back_btn,
                                         callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "slow":
                rows = await DB.get_slow_queries(20)
                if not rows:
                    await safe_edit(
                        query,
                        await _trans('slow_queries_title', lang, "🐌") + "\n\n"
                        + await _trans('no_slow_queries', lang, "📭"),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        parse_mode='HTML', bot=context.bot)
                    return
                text = await _trans('slow_queries_title', lang, "🐌") + "\n"
                text += "━━━━━━━━━━━━━━━━━━━━━━\n\n"
                for i, q in enumerate(rows[:10], 1):
                    elapsed = q.get('elapsed', 0)
                    color = "🔴" if elapsed >= 3 else ("🟡" if elapsed >= 1.5 else "🟢")
                    caller_file = q.get('caller_file', '?') or '?'
                    caller_line = q.get('caller_line', 0)
                    caller_func = q.get('caller_func', '?') or '?'
                    conn_type = q.get('conn_type', '?')
                    params_count = q.get('params_count', 0)
                    qs = _html.escape(str(q.get('query', ''))[:70])
                    text += f"{i}. {color} <code>{elapsed:.2f}s</code>\n"
                    text += f"   <i>{qs}</i>\n"
                    text += (await _trans('slow_queries_source', lang, "📍") + " "
                             + f"<code>{_html.escape(caller_file)}:{caller_line}</code>\n")
                    text += (await _trans('slow_queries_function', lang, "🔧") + " "
                             + f"<code>{_html.escape(caller_func)}</code>\n")
                    extra = []
                    if conn_type and conn_type != '?':
                        extra.append(f"🗄️ {conn_type}")
                    if params_count:
                        extra.append(f"🔢 {params_count}")
                    if extra:
                        text += f"   {' | '.join(extra)}\n"
                    text += "\n"
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('refresh_btn', lang, "🔄"),
                        callback_data="analytics_slow")],
                    [InlineKeyboardButton(back_btn,
                                          callback_data="admin_analytics")]])
                await safe_edit(query, text, reply_markup=kb,
                                parse_mode='HTML', bot=context.bot)
                return

            if action == "export":
                await safe_edit(query,
                    await _trans('export_starting', lang, "⏳"), bot=context.bot)
                try:
                    file_path = await CallbackHandlers._generate_excel_report()
                    with open(file_path, 'rb') as f:
                        await context.bot.send_document(
                            chat_id=user_id, document=f,
                            filename=os.path.basename(file_path),
                            caption=await _trans('report_caption', lang, "📊"),
                            parse_mode='HTML')
                    try:
                        os.remove(file_path)
                    except OSError:
                        pass
                    await safe_edit(
                        query,
                        await _trans('export_success', lang, "✅"),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        bot=context.bot)
                except ImportError:
                    await safe_edit(
                        query,
                        await _trans('export_no_openpyxl', lang, "❌"),
                        parse_mode='HTML',
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        bot=context.bot)
                except Exception as e:
                    logger.error(f"❌ export: {e}", exc_info=True)
                    await safe_edit(
                        query,
                        _fmt(await _trans('export_failed', lang, "❌ {error}"),
                             error=str(e)[:80]),
                        reply_markup=InlineKeyboardMarkup([[
                            InlineKeyboardButton(back_btn,
                                                 callback_data="admin_analytics")]]),
                        bot=context.bot)
                return

            await safe_edit(query,
                await _trans('unknown_report', lang, "⚠️"), bot=context.bot)

        except Exception as e:
            logger.error(f"❌ _handle_analytics({action}): {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _generate_excel_report() -> str:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="4472C4",
                                  end_color="4472C4", fill_type="solid")

        ws1 = wb.active
        ws1.title = "Users"
        ws1['A1'] = "Date"
        ws1['B1'] = "Count"
        ws1['A1'].font = header_font
        ws1['B1'].font = header_font
        ws1['A1'].fill = header_fill
        ws1['B1'].fill = header_fill
        growth = await DB.get_user_growth(30)
        for i, r in enumerate(growth, start=2):
            ws1[f'A{i}'] = r['date']
            ws1[f'B{i}'] = r['count']

        ws2 = wb.create_sheet("Channels")
        headers = ["#", "Name", "Total", "Published", "Failed",
                   "Success %", "Completion %"]
        for i, h in enumerate(headers, 1):
            c = ws2.cell(row=1, column=i, value=h)
            c.font = header_font
            c.fill = header_fill
        top = await DB.get_top_channels(20)
        for i, ch in enumerate(top, start=2):
            ws2.cell(row=i, column=1, value=i - 1)
            ws2.cell(row=i, column=2, value=ch['name'])
            ws2.cell(row=i, column=3, value=ch['total'])
            ws2.cell(row=i, column=4, value=ch['published'])
            ws2.cell(row=i, column=5, value=ch['failed'])
            ws2.cell(row=i, column=6, value=ch.get('success_rate', 100))
            ws2.cell(row=i, column=7, value=ch.get('completion_rate', 0))

        ws3 = wb.create_sheet("Subscriptions")
        ws3['A1'] = "Month"
        ws3['B1'] = "Count"
        ws3['A1'].font = header_font
        ws3['B1'].font = header_font
        ws3['A1'].fill = header_fill
        ws3['B1'].fill = header_fill
        subs = await DB.get_subscription_rate(12)
        for i, r in enumerate(subs, start=2):
            ws3[f'A{i}'] = r['month']
            ws3[f'B{i}'] = r['count']

        ws4 = wb.create_sheet("Stats")
        stats = await DB.get_general_stats()
        pub_stats = await DB.get_publish_stats()
        pool = await DB.get_pool_live()
        rows = [
            ("Users", stats.get('users', 0)),
            ("Channels", stats.get('channels', 0)),
            ("Groups", stats.get('groups', 0)),
            ("Posts", stats.get('posts', 0)),
            ("Published", stats.get('published', 0)),
            ("Invoices", stats.get('invoices', 0)),
            ("Tickets", stats.get('tickets', 0)),
            ("DB size KB", stats.get('db_size_kb', 0)),
            ("", ""),
            ("Avg posts/channel", pub_stats['avg_posts_per_channel']),
            ("Success %", f"{pub_stats.get('success_rate', 100)}%"),
            ("Completion %", f"{pub_stats.get('completion_rate', 0)}%"),
            ("", ""),
            ("Pool available", pool.get('available', False)),
        ]
        for i, (k, v) in enumerate(rows, start=1):
            ws4[f'A{i}'] = k
            ws4[f'B{i}'] = v
            if i == 1:
                ws4[f'A{i}'].font = header_font
                ws4[f'B{i}'].font = header_font

        import tempfile
        ts = TimeUtils.utc_now().strftime('%Y%m%d_%H%M%S')
        file_path = os.path.join(tempfile.gettempdir(), f"analytics_{ts}.xlsx")
        wb.save(file_path)
        return file_path

    # ═════════════════════════════════════════════════════════════
    # Auto replies
    # =================================================================

    @staticmethod
    async def _handle_auto_reply(update, context, query, user_id, lang=None):
        if not lang:
            lang = await DB.get_user_language(user_id) or 'ar'
        data = query.data
        parts = data.split(":")
        raw_action = parts[0]
        if raw_action.startswith("sec_"):
            raw_action = raw_action[4:]
        if raw_action.startswith("auto_reply_"):
            action = raw_action[len("auto_reply_"):]
        else:
            action = raw_action

        chat_id = None
        if len(parts) >= 2:
            p = parts[1].strip()
            if p.lstrip('-').isdigit():
                chat_id = int(p)
        if chat_id is None:
            stored = context.user_data.get('auto_chat')
            if stored is not None:
                try:
                    chat_id = int(stored)
                except (TypeError, ValueError):
                    chat_id = None
        if chat_id is None:
            await safe_edit(query,
                await _trans('group_not_specified', lang, "❌"), bot=context.bot)
            return

        if chat_id == -1:
            if not CONFIG.is_developer(user_id):
                await safe_edit(query,
                    await _trans('unauthorized', lang, "❌"), bot=context.bot)
                return
        else:
            if not await is_authorized_in_group(context.bot, chat_id, user_id):
                await safe_edit(query,
                    await _trans('no_permission', lang, "❌"), bot=context.bot)
                return

        try:
            if action == "menu":
                context.user_data['auto_chat'] = chat_id
                kb = KeyboardFactory.build("auto_reply",
                                            chat_id=chat_id, lang=lang)
                await safe_edit(query,
                    await _trans('auto_reply_title', lang, "🤖"),
                    reply_markup=kb, bot=context.bot)
                return

            if action == "toggle":
                cache_key = f"ars_{chat_id}"
                settings = context.user_data.get(cache_key)
                if settings is None:
                    settings = await DB.get_auto_reply_settings(chat_id) or {}
                    if not isinstance(settings, dict):
                        settings = _row_to_dict(settings) or {}
                new_status = not settings.get('enabled', False)
                await DB.update_auto_reply_settings(chat_id, enabled=new_status)
                settings['enabled'] = new_status
                context.user_data[cache_key] = settings
                kb = KeyboardFactory.build("auto_reply",
                                            chat_id=chat_id, lang=lang)
                status = (await _trans('auto_reply_status_enabled', lang, "✅")
                          if new_status
                          else await _trans('auto_reply_status_disabled', lang, "❌"))
                admins = (await _trans('auto_reply_admins_yes', lang, "✅")
                          if settings.get('only_admins')
                          else await _trans('auto_reply_admins_no', lang, "❌"))
                text = _fmt(await _trans('auto_reply_full', lang,
                                          "🤖\n{status}\n{admins}"),
                            status=status, admins=admins)
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if action == "admins":
                cache_key = f"ars_{chat_id}"
                settings = context.user_data.get(cache_key)
                if settings is None:
                    settings = await DB.get_auto_reply_settings(chat_id) or {}
                    if not isinstance(settings, dict):
                        settings = _row_to_dict(settings) or {}
                new_status = not settings.get('only_admins', 0)
                await DB.update_auto_reply_settings(chat_id, only_admins=new_status)
                settings['only_admins'] = new_status
                context.user_data[cache_key] = settings
                kb = KeyboardFactory.build("auto_reply",
                                            chat_id=chat_id, lang=lang)
                status = (await _trans('auto_reply_status_enabled', lang, "✅")
                          if settings.get('enabled')
                          else await _trans('auto_reply_status_disabled', lang, "❌"))
                admins = (await _trans('auto_reply_admins_yes', lang, "✅")
                          if new_status
                          else await _trans('auto_reply_admins_no', lang, "❌"))
                text = _fmt(await _trans('auto_reply_full', lang,
                                          "🤖\n{status}\n{admins}"),
                            status=status, admins=admins)
                await safe_edit(query, text, reply_markup=kb, bot=context.bot)
                return

            if action == "add":
                StateManager.set(user_id, UserState.WAIT_AUTO_KEY)
                context.user_data['auto_chat'] = chat_id
                await safe_edit(query,
                    await _trans('send_keyword_prompt', lang, "📝"),
                    bot=context.bot)
                return

            if action == "del":
                StateManager.set(user_id, UserState.WAIT_AUTO_DEL)
                context.user_data['auto_chat'] = chat_id
                await safe_edit(query,
                    await _trans('send_keyword_delete_prompt', lang, "🗑️"),
                    bot=context.bot)
                return

            if action == "reset":
                kb = InlineKeyboardMarkup([
                    [InlineKeyboardButton(
                        await _trans('auto_reply_reset_btn', lang, "🗑️"),
                        callback_data=f"auto_reply_reset_confirm:{chat_id}")],
                    [InlineKeyboardButton(
                        await _trans('cancel_btn', lang, "❌"),
                        callback_data=f"auto_reply_menu:{chat_id}")],
                ])
                await safe_edit(query,
                    await _trans('auto_reply_reset_confirm', lang, "⚠️"),
                    reply_markup=kb, parse_mode='HTML', bot=context.bot)
                return

            if action == "reset_confirm":
                await DB.reset_auto_replies(chat_id)
                context.user_data.pop(f"ars_{chat_id}", None)
                await safe_edit(query,
                    await _trans('auto_reply_reset_success', lang, "✅"),
                    bot=context.bot)
                return

            if action == "list":
                rows = await DB.fetchall(
                    "SELECT keyword FROM auto_replies "
                    "WHERE chat_id=? LIMIT 20", (chat_id,))
                if rows:
                    lines = []
                    for r in rows:
                        rd = _row_to_dict(r) or {}
                        lines.append(f"• {_safe_str(rd.get('keyword'))}")
                    text = await _trans('auto_reply_list_title', lang, "📋") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_auto_replies', lang, "📭")
                await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=f"auto_reply_menu:{chat_id}")]]),
                    bot=context.bot)
                return

            if action == "stats":
                stats = await DB.get_auto_reply_stats(chat_id, 20)
                if stats:
                    text = await _trans('auto_reply_stats_title', lang, "📊") + "\n\n"
                    for s in stats:
                        sd = _row_to_dict(s) or {}
                        source = (await _trans('auto_reply_source_global', lang, "🌐")
                                  if sd.get('source') == 'global'
                                  else await _trans('auto_reply_source_group', lang, "👥"))
                        text += _fmt(await _trans('auto_reply_stats_line', lang,
                                                  "• {keyword} ({source}): {count}"),
                                     keyword=_safe_str(sd.get('keyword')),
                                     source=source,
                                     count=sd.get('usage_count', 0)) + "\n"
                else:
                    text = await _trans('no_auto_replies_stats', lang, "📭")
                await safe_edit(query, text, reply_markup=InlineKeyboardMarkup(
                    [[InlineKeyboardButton(
                        KeyboardFactory.get_text("back", lang),
                        callback_data=f"auto_reply_menu:{chat_id}")]]),
                    bot=context.bot)
                return

            await safe_edit(query,
                await _trans('unknown_action', lang, "⚠️"), bot=context.bot)
        except Exception as e:
            logger.error(f"auto_reply error: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Schedule
    # =================================================================

    @staticmethod
    async def _handle_schedule(update, context, query, user_id):
        lang = await DB.get_user_language(user_id) or 'ar'
        data = query.data
        parts = data.split(":")
        if len(parts) < 2:
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        action = parts[0].replace("sched_", "")
        try:
            ch_id = int(parts[1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        if not await _is_channel_owner(user_id, ch_id):
            await safe_edit(query,
                await _trans('not_owner', lang, "❌"), bot=context.bot)
            return
        try:
            if action == "open":
                await CallbackHandlers._show_schedule_menu(
                    update, context, query, ch_id, user_id)
                return
            if action == "min":
                StateManager.set(user_id, UserState.WAIT_MIN)
                context.user_data['schedule_ch'] = ch_id
                await safe_edit(query,
                    await _trans('send_minutes_btn', lang, "📅"), bot=context.bot)
                return
            if action == "hour":
                StateManager.set(user_id, UserState.WAIT_HOUR)
                context.user_data['schedule_ch'] = ch_id
                await safe_edit(query,
                    await _trans('send_hours_btn', lang, "📅"), bot=context.bot)
                return
            if action == "day":
                StateManager.set(user_id, UserState.WAIT_DAY)
                context.user_data['schedule_ch'] = ch_id
                await safe_edit(query,
                    await _trans('send_days_btn', lang, "📅"), bot=context.bot)
                return
            if action == "time":
                StateManager.set(user_id, UserState.WAIT_PUB_TIME)
                context.user_data['schedule_ch'] = ch_id
                await safe_edit(query,
                    await _trans('send_time_btn', lang, "🕐"), bot=context.bot)
                return
            await safe_edit(query,
                await _trans('unknown_action', lang, "⚠️"), bot=context.bot)
        except Exception as e:
            logger.error(f"schedule error: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    @staticmethod
    async def _show_schedule_menu(update, context, query, ch_id, user_id):
        lang = await DB.get_user_language(user_id) or 'ar'
        kb = KeyboardFactory.build("channel_settings", chat_id=ch_id, lang=lang)
        await safe_edit(query,
            await _trans('channel_schedule_title', lang, "📅"),
            reply_markup=kb, bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Advanced actions
    # =================================================================

    @staticmethod
    async def _handle_advanced_actions(update, context, query, user_id):
        lang = await DB.get_user_language(user_id) or 'ar'
        data = query.data
        parts = data.split(":")
        if len(parts) < 2:
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        prefix = parts[0]
        action = prefix.replace("act_", "").replace("pen_", "").replace("ban_", "")
        try:
            chat_id = int(parts[1])
        except (ValueError, IndexError):
            await safe_edit(query,
                await _trans('invalid_data', lang, "❌"), bot=context.bot)
            return
        if chat_id == -1 and (prefix.startswith("act_") or prefix.startswith("pen_")):
            await safe_edit(query,
                await _trans('invalid_id', lang, "❌"), bot=context.bot)
            return
        if chat_id != -1:
            if not await is_authorized_in_group(context.bot, chat_id, user_id):
                await safe_edit(query,
                    await _trans('no_permission', lang, "❌"), bot=context.bot)
                return
        else:
            if not CONFIG.is_developer(user_id):
                await safe_edit(query,
                    await _trans('unauthorized', lang, "❌"), bot=context.bot)
                return

        try:
            if prefix.startswith("ban_"):
                if action == "add":
                    state = (UserState.WAIT_GROUP_BAN if chat_id != -1
                             else UserState.WAIT_GLOBAL_BAN)
                    StateManager.set(user_id, state)
                    context.user_data['ban_chat'] = chat_id
                    await safe_edit(query,
                        await _trans('send_keyword_prompt', lang, "📝"),
                        bot=context.bot)
                    return
                if action == "list":
                    words = await DB.get_banned_words(chat_id)
                    if words:
                        text = (await _trans('words_list_title_full', lang, "🚫")
                                + "\n\n"
                                + "\n".join(f"• {w}" for w in words[:50]))
                    else:
                        text = await _trans('no_data', lang, "📭")
                    await safe_edit(query, text, bot=context.bot)
                    return
                if action == "rem":
                    state = (UserState.WAIT_REM_GROUP_BAN if chat_id != -1
                             else UserState.WAIT_REM_GLOBAL_BAN)
                    StateManager.set(user_id, state)
                    context.user_data['ban_chat'] = chat_id
                    await safe_edit(query,
                        await _trans('send_keyword_delete_prompt', lang, "🗑️"),
                        bot=context.bot)
                    return
                return

            if prefix.startswith("act_"):
                user_actions = {
                    "ban": (UserState.WAIT_BAN, "send_user_id_ban", "🚫"),
                    "mute": (UserState.WAIT_MUTE, "send_user_id_mute", "🔇"),
                    "warn": (UserState.WAIT_WARN, "send_user_id_warn", "⚠️"),
                    "kick": (UserState.WAIT_KICK, "send_user_id_kick", "👢"),
                    "restrict": (UserState.WAIT_RESTRICT, "send_user_id_restrict", "🔒"),
                    "unban": (UserState.WAIT_UNBAN, "send_user_id_unban", "🔓"),
                    "pin": (UserState.WAIT_PIN, "pin_prompt_full", "📌"),
                }
                if action in user_actions:
                    state, prompt_key, prompt_default = user_actions[action]
                    StateManager.set(user_id, state)
                    context.user_data['adv_chat'] = chat_id
                    await safe_edit(query,
                        await _trans(prompt_key, lang, prompt_default),
                        bot=context.bot)
                    return
                if action == "log":
                    await CallbackHandlers._show_admin_logs(
                        update, context, query, chat_id, lang)
                    StateManager.clear(user_id)
                    return
                return

            if prefix.startswith("pen_"):
                penalty_types = {'ban', 'mute', 'kick', 'restrict', 'none'}
                if action in penalty_types:
                    await DB.update_security_settings(chat_id, auto_penalty=action)
                    await CallbackHandlers._invalidate_security_settings_cache(chat_id)
                    await CallbackHandlers._refresh_security_view(
                        query, context, chat_id, lang)
                    return
                return
        except Exception as e:
            logger.error(f"advanced actions: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Panel
    # =================================================================

    @staticmethod
    async def _handle_panel(update, context, query, user_id, data, lang='ar'):
        if not update.effective_chat:
            await safe_edit(query,
                await _trans('cannot_determine_group', lang, "❌"),
                bot=context.bot)
            return
        chat_id = update.effective_chat.id
        if not await is_authorized_in_group(context.bot, chat_id, user_id):
            await safe_edit(query,
                await _trans('no_permission', lang, "❌"), bot=context.bot)
            return
        try:
            if data == "panel_lock":
                await context.bot.set_chat_permissions(
                    chat_id,
                    permissions=ChatPermissions(
                        can_send_messages=False, can_send_audios=False,
                        can_send_documents=False,
                        can_send_photos=False, can_send_videos=False,
                        can_send_video_notes=False,
                        can_send_voice_notes=False, can_send_polls=False,
                        can_send_other_messages=False,
                        can_add_web_page_previews=False, can_change_info=False,
                        can_invite_users=False, can_pin_messages=False))
                await safe_edit(query,
                    await _trans('group_locked_full', lang, "🔒"), bot=context.bot)
                return
            if data == "panel_unlock":
                await context.bot.set_chat_permissions(
                    chat_id,
                    permissions=ChatPermissions(
                        can_send_messages=True, can_send_audios=True,
                        can_send_documents=True,
                        can_send_photos=True, can_send_videos=True,
                        can_send_video_notes=True,
                        can_send_voice_notes=True, can_send_polls=True,
                        can_send_other_messages=True,
                        can_add_web_page_previews=True, can_change_info=True,
                        can_invite_users=True, can_pin_messages=True))
                await safe_edit(query,
                    await _trans('group_unlocked_full', lang, "🔓"),
                    bot=context.bot)
                return
            if data == "panel_close":
                StateManager.clear(user_id)
                _clear_context_keys(context)
                await safe_delete_message(query)
                return
        except Exception as e:
            logger.error(f"panel: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Contests
    # =================================================================

    @staticmethod
    async def _handle_contests(update, context, query, user_id):
        lang = await DB.get_user_language(user_id) or 'ar'
        data = query.data
        declare_sel = getattr(CB, 'DECLARE_WINNER_SEL', 'declare_winner_sel')
        contest_join = getattr(CB, 'CONTEST_JOIN', 'contest_join')
        try:
            if data.startswith(contest_join + ":"):
                cid = _coerce_int(data.split(":")[-1])
                if cid <= 0:
                    await safe_edit(query,
                        await _trans('invalid_data', lang, "❌"), bot=context.bot)
                    return
                contest = await DB.get_contest_by_id(cid)
                cd = _row_to_dict(contest) or {}
                if not cd or cd.get('status') != 'active':
                    await safe_edit(query,
                        await _trans('no_contests', lang, "❌"), bot=context.bot)
                    StateManager.clear(user_id)
                    return
                already_joined = await DB.check_contest_joined(cid, user_id)
                if already_joined:
                    await safe_edit(query,
                        await _trans('join_contest', lang, "❌") + " ✅",
                        bot=context.bot)
                    return
                StateManager.set(user_id, UserState.WAIT_CONTEST_ANSWER)
                context.user_data['contest_join'] = cid
                await safe_edit(query,
                    await _trans('send_answer_btn', lang, "📝"), bot=context.bot)
                return

            if data == CB.CONTEST_WINNERS:
                winners = await DB.get_contest_winners(10)
                if winners:
                    lines = []
                    for w in winners:
                        wd = _row_to_dict(w) or {}
                        lines.append(f"• {_safe_str(wd.get('title'))} - "
                                     f"{_safe_str(wd.get('winner_id'))}")
                    text = await _trans('winners_title', lang, "🏆") + "\n\n" + "\n".join(lines)
                else:
                    text = await _trans('no_winners', lang, "📭")
                await safe_edit(query, text, bot=context.bot)
                StateManager.clear(user_id)
                return

            if data.startswith(declare_sel + ":"):
                if not CONFIG.is_developer(user_id):
                    await safe_edit(query,
                        await _trans('unauthorized', lang, "❌"), bot=context.bot)
                    return
                cid = _coerce_int(data.split(":")[-1])
                if cid <= 0:
                    await safe_edit(query,
                        await _trans('invalid_data', lang, "❌"), bot=context.bot)
                    return
                participants = await DB.fetchall(
                    "SELECT user_id FROM contest_participants "
                    "WHERE contest_id=?", (cid,))
                if not participants:
                    await safe_edit(query,
                        await _trans('no_participants_full', lang, "❌"),
                        bot=context.bot)
                    return
                user_ids = []
                for p in participants:
                    pd = _row_to_dict(p) or {}
                    uid_val = pd.get('user_id')
                    if uid_val is not None:
                        user_ids.append(uid_val)
                if not user_ids:
                    await safe_edit(query,
                        await _trans('no_participants_full', lang, "❌"),
                        bot=context.bot)
                    return
                winner_id = random.choice(user_ids)
                if await DB.declare_winner(cid, winner_id):
                    msg = _fmt(await _trans('winner_announced', lang, "✅ {winner_id}"),
                               winner_id=winner_id)
                    await safe_edit(query, msg, bot=context.bot)
                    try:
                        congrats = await _trans('congrats_winner_full', lang, "🎉")
                        await context.bot.send_message(winner_id, congrats)
                    except Exception:
                        pass
                else:
                    await safe_edit(query,
                        await _trans('declare_failed', lang, "❌"), bot=context.bot)
                return
        except Exception as e:
            logger.error(f"contests: {e}", exc_info=True)
            await safe_edit(query,
                await _trans('error_occurred', lang, "❌"), bot=context.bot)

    # ═════════════════════════════════════════════════════════════
    # Backup
    # =================================================================

    @staticmethod
    async def _do_backup(context, user_id, lang='ar'):
        try:
            PATHS.BACKUPS.mkdir(parents=True, exist_ok=True)
            timestamp = TimeUtils.mecca_now().strftime('%Y%m%d_%H%M%S_%f')
            backup_file = PATHS.BACKUPS / f"backup_{timestamp}.db"
            success = await DB.backup_database(backup_file)
            if not success:
                await safe_send(context.bot, user_id,
                                await _trans('backup_failed_full', lang, "❌"))
                return
            try:
                await DB.set_setting('last_backup', TimeUtils.sql_iso())
            except Exception:
                pass
            try:
                size = backup_file.stat().st_size
            except Exception:
                size = 0
            if size > MAX_TG_FILE_SIZE:
                size_mb = size / (1024 * 1024)
                msg = _fmt(await _trans('backup_too_large', lang,
                                        "⚠️ {size}MB {path}"),
                           size=f"{size_mb:.1f}", path=str(backup_file))
                await safe_send(context.bot, user_id, msg)
                return

            def _safe_mtime(p):
                try:
                    return p.stat().st_mtime
                except (OSError, FileNotFoundError):
                    return 0

            backups = sorted(PATHS.BACKUPS.glob("backup_*.db"),
                             key=_safe_mtime, reverse=True)
            others = [b for b in backups if b != backup_file]
            for old in others[MAX_BACKUPS - 1:]:
                try:
                    old.unlink(missing_ok=True)
                except OSError:
                    pass

            with open(backup_file, 'rb') as f:
                await context.bot.send_document(
                    chat_id=user_id, document=f, filename=backup_file.name)
        except Exception as e:
            logger.error(f"❌ backup: {e}", exc_info=True)
            try:
                msg = _fmt(await _trans('backup_error_full', lang, "❌ {error}"),
                           error=str(e)[:100])
                await safe_send(context.bot, user_id, msg)
            except Exception:
                pass


__all__ = [
    "CallbackHandlers",
    "_invalidate_sec_auth_cache",
    "_invalidate_after_channel_change",
    "_set_sec_chat",
    "_format_channel_rate_line",
    "_md_to_html",
    "_get_log_channel_menu_data",
    "_invalidate_log_channel_menu_cache",
    "_ANALYTICS_ALIASES",
    "ACTIVE_TASKS",
]